import json
import base64
import hashlib
import io
import os
import subprocess
import sys
import tempfile
import threading
import unittest
import zipfile
import shutil
import uuid
import openpyxl
from types import SimpleNamespace
from pathlib import Path
from unittest.mock import patch

from werkzeug.serving import make_server

from remote_storage import (
    RemoteJsonStorage, RemoteStorageConflict, RemoteStorageError,
    configured_remote_storage,
)
from storage_server import (
    StorageServiceError,
    create_app,
    create_backup,
    configured_storage_tokens,
    inspect_backup,
    restore_backup,
    version_for,
)


class RemoteStorageServerTests(unittest.TestCase):
    def setUp(self):
        # Optional inherited-permission directory for restricted Windows runners.
        # Normal CI uses TemporaryDirectory; neither branch touches real data.
        test_root = os.getenv('TEMLI_TEST_TMP_ROOT')
        if test_root:
            directory = Path(test_root).resolve() / ('test-' + uuid.uuid4().hex)
            directory.mkdir()
            self.temp = SimpleNamespace(name=str(directory))
            self.addCleanup(shutil.rmtree, directory)
        else:
            self.temp = tempfile.TemporaryDirectory()
            self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name) / "temli"
        self.backups = Path(self.temp.name) / "backups"
        self.token = "test-token-" + "x" * 40
        self.app = create_app(
            self.root,
            self.token,
            initialize=True,
            backups=self.backups,
        )
        self.client = self.app.test_client()
        self.headers = {"Authorization": "Bearer " + self.token}

    def post(self, path, body, headers=None):
        return self.client.post(
            path,
            json=body,
            headers=self.headers if headers is None else headers,
        )

    def test_initializes_empty_storage_and_round_trips_json(self):
        response = self.post("/v1/json/read", {"path": "teacher_data/101/schedule.json"})
        self.assertEqual(response.status_code, 200)
        self.assertFalse(response.get_json()["exists"])

        response = self.post("/v1/json/write", {
            "path": "teacher_data/101/schedule.json",
            "data": {"2026-09-20": [{"student": "Тест"}]},
        })
        self.assertEqual(response.status_code, 200)
        response = self.post("/v1/json/read", {"path": "teacher_data/101/schedule.json"})
        self.assertEqual(
            response.get_json()["data"]["2026-09-20"][0]["student"],
            "Тест",
        )

    def test_authenticated_status_reports_backup_without_tenant_data(self):
        self.assertEqual(self.post("/v1/status", {}, {}).status_code, 401)
        empty = self.post("/v1/status", {})
        self.assertEqual(empty.status_code, 200)
        self.assertEqual(empty.get_json()["backup"]["count"], 0)
        self.assertFalse(empty.get_json()["backup"]["latest_verified"])
        create_backup(self.root, self.backups, reason="manual")
        ready = self.post("/v1/status", {}).get_json()
        self.assertEqual(ready["status"], "ok")
        self.assertEqual(ready["backup"]["count"], 1)
        self.assertIsInstance(ready["backup"]["latest_age_seconds"], int)
        self.assertTrue(ready["backup"]["latest_verified"])
        self.assertNotIn("data", ready)

    def test_authenticated_status_rejects_a_corrupt_latest_backup_safely(self):
        archive, _ = create_backup(self.root, self.backups, reason="manual")
        archive.write_bytes(b"not a zip archive")

        response = self.post("/v1/status", {})

        self.assertEqual(response.status_code, 503)
        self.assertEqual(response.get_json(), {
            "status": "error",
            "code": "backup_verification_failed",
        })
        self.assertNotIn(archive.name, response.get_data(as_text=True))

    def test_authenticated_status_caches_verification_of_unchanged_backup(self):
        create_backup(self.root, self.backups, reason="manual")
        with patch("storage_server.inspect_backup", wraps=inspect_backup) as verify:
            self.assertEqual(self.post("/v1/status", {}).status_code, 200)
            self.assertEqual(self.post("/v1/status", {}).status_code, 200)
        self.assertEqual(verify.call_count, 1)

    def test_batch_json_read_and_real_client_versions(self):
        self.post("/v1/json/write", {"path": "schedule.json", "data": {"week": 1}})
        response = self.post("/v1/json/read-batch", {
            "paths": ["schedule.json", "students.json"],
        })
        self.assertEqual(response.status_code, 200)
        items = response.get_json()["items"]
        self.assertTrue(items["schedule.json"]["exists"])
        self.assertEqual(items["students.json"]["data"], {})

        server = make_server('127.0.0.1', 0, self.app)
        thread = threading.Thread(target=server.serve_forever, daemon=True)
        thread.start()
        try:
            with patch.dict(os.environ, TEMLI_STORAGE_ALLOW_HTTP='true'):
                remote = RemoteJsonStorage('http://127.0.0.1:' + str(server.server_port), self.token)
                remote_status = remote.status()
                self.assertEqual(remote_status["service"], "temli-storage")
                self.assertEqual(remote_status["backup"]["count"], 0)
                values = remote.read_json_batch({"schedule.json": {}, "students.json": {"empty": True}})
                self.assertEqual(values["schedule.json"], {"week": 1})
                self.assertEqual(values["students.json"], {})
                remote.write_json("students.json", {"created": True})
                self.assertEqual(remote.read_json("students.json", {}), {"created": True})
        finally:
            server.shutdown()
            server.server_close()
            thread.join(timeout=5)

        for paths in ([], ["schedule.json"] * 2, ["schedule.json"] * 17):
            self.assertEqual(self.post("/v1/json/read-batch", {"paths": paths}).status_code, 400)

    def test_binary_roundtrip_conflict_and_auth(self):
        path = 'teacher_data/101/book.xlsx'
        raw = b'workbook-bytes-for-storage-test'
        self.assertEqual(self.post('/v1/files/read', {'path': path}, {}).status_code, 401)
        self.assertFalse(self.post('/v1/files/read', {'path': path}).get_json()['exists'])
        body = {'path': path, 'data': base64.b64encode(raw).decode(), 'expected_version': None}
        created = self.post('/v1/files/write', body)
        self.assertEqual(created.status_code, 200)
        result = self.post('/v1/files/read', {'path': path}).get_json()
        self.assertEqual(base64.b64decode(result['data']), raw)
        self.assertEqual(result['version'], hashlib.sha256(raw).hexdigest())
        self.assertEqual(self.post('/v1/files/write', body).status_code, 409)
        body['expected_version'] = result['version']
        body['data'] = base64.b64encode(b'updated').decode()
        self.assertEqual(self.post('/v1/files/write', body).status_code, 200)
        self.assertEqual((self.root / (path + '.bak')).read_bytes(), raw)

        deleted = self.post('/v1/files/delete', {
            'path': path,
            'expected_version': hashlib.sha256(b'updated').hexdigest(),
        })
        self.assertEqual(deleted.status_code, 200)
        self.assertTrue(deleted.get_json()['existed'])
        self.assertFalse((self.root / path).exists())
        self.assertEqual((self.root / (path + '.bak')).read_bytes(), b'updated')
        self.assertEqual(self.post('/v1/files/delete', {
            'path': path,
            'expected_version': hashlib.sha256(b'updated').hexdigest(),
        }).status_code, 409)

    def test_binary_paths_and_invalid_requests_are_rejected(self):
        for path in ('../book.xlsx', '/book.xlsx', 'C:/book.xlsx',
                     'teacher_data/../book.xlsx', 'teacher_data/bad/book.xlsx',
                     'receipt_assets/other.png', 'bot.py', 'students.json',
                     'receipt_assets//logo.png'):
            with self.subTest(path=path):
                self.assertEqual(self.post('/v1/files/read', {'path': path}).status_code, 400)
        self.assertEqual(self.post('/v1/files/write', []).status_code, 400)
        self.assertEqual(self.post('/v1/files/write', {'path': 'book.xlsx', 'data': '%%%','expected_version': None}).status_code, 400)
        self.assertEqual(self.post('/v1/files/write', {'path': 'book.xlsx', 'data': 'eA=='}).status_code, 400)

    def test_binary_backups_restore_bytes_and_read_legacy_archive(self):
        legacy, _ = create_backup(self.root, self.backups)
        # Existing schema-1 JSON archives must remain readable.
        old_path = self.backups / 'legacy.zip'
        with zipfile.ZipFile(legacy) as source, zipfile.ZipFile(old_path, 'w') as target:
            for name in source.namelist():
                raw = source.read(name)
                if name == 'manifest.json':
                    value = json.loads(raw)
                    value['schema'] = 1
                    raw = json.dumps(value).encode()
                target.writestr(name, raw)
        self.assertEqual(inspect_backup(old_path)[0]['schema'], 1)
        files = {'book.xlsx': b'workbook', 'receipt_assets/logo.png': b'image',
                 'teacher_data/101/receipts/check_1.pdf': b'%PDF-test'}
        for path, raw in files.items():
            response = self.post('/v1/files/write', {'path': path, 'data': base64.b64encode(raw).decode(), 'expected_version': None})
            self.assertEqual(response.status_code, 200)
        archive, manifest = create_backup(self.root, self.backups)
        self.assertEqual(manifest['schema'], 2)
        _, payload = inspect_backup(archive)
        for path, raw in files.items():
            self.assertEqual(payload[path], raw)
        (self.root / 'book.xlsx').write_bytes(b'changed')
        restore_backup(self.root, archive, self.backups)
        self.assertEqual((self.root / 'book.xlsx').read_bytes(), b'workbook')

    def test_real_client_binary_roundtrip_and_conflict(self):
        server = make_server('127.0.0.1', 0, self.app)
        thread = threading.Thread(target=server.serve_forever, daemon=True)
        thread.start()
        try:
            with patch.dict(os.environ, TEMLI_STORAGE_ALLOW_HTTP='true'):
                first = RemoteJsonStorage('http://127.0.0.1:' + str(server.server_port), self.token)
                second = RemoteJsonStorage(first.base_url, self.token)
                with self.assertRaises(RemoteStorageError):
                    first.write_file('book.xlsx', b'unsafe-blind-write')
                self.assertIsNone(first.read_file('book.xlsx'))
                self.assertIsNone(second.read_file('book.xlsx'))
                first.write_file('book.xlsx', b'original')
                with self.assertRaises(RemoteStorageConflict):
                    second.write_file('book.xlsx', b'lost-update')
                self.assertEqual(second.read_file('book.xlsx'), b'original')
                self.assertTrue(second.delete_file('book.xlsx'))
                self.assertIsNone(second.read_file('book.xlsx'))
        finally:
            server.shutdown()
            server.server_close()
            thread.join(timeout=5)

    def test_atomic_payment_transaction_and_idempotent_server_retry(self):
        server = make_server('127.0.0.1', 0, self.app)
        thread = threading.Thread(target=server.serve_forever, daemon=True)
        thread.start()
        try:
            with patch.dict(os.environ, TEMLI_STORAGE_ALLOW_HTTP='true'):
                remote = RemoteJsonStorage('http://127.0.0.1:' + str(server.server_port), self.token)
                self.assertEqual(remote.read_json('schedule.json', {}), {})
                self.assertEqual(remote.read_json('payments.json', {}), {})
                self.assertIsNone(remote.read_file('book.xlsx'))
                transaction_id = 'payment-' + 'a' * 32
                remote.commit_payment_transaction(
                    {'schedule.json': {'lesson': {'paid': True}},
                     'payments.json': {'tx': {'amount': 1000}}},
                    {'book.xlsx': b'xlsx-payment-book'},
                    transaction_id=transaction_id,
                )
                self.assertTrue(remote.read_json('schedule.json', {})['lesson']['paid'])
                self.assertEqual(remote.read_json('payments.json', {})['tx']['amount'], 1000)
                self.assertEqual(remote.read_file('book.xlsx'), b'xlsx-payment-book')

                archive, _manifest = create_backup(self.root, self.backups)
                _checked, payload = inspect_backup(archive)
                self.assertEqual(payload['book.xlsx'], b'xlsx-payment-book')
                self.assertIn('payments.json', payload)
                self.assertFalse(any('.payment-transactions' in path for path in payload))

            # A lost HTTP response can be retried with the exact same payload.
            items = [
                {'path': 'schedule.json', 'kind': 'json', 'data': {'next': True},
                 'expected_version': version_for({'lesson': {'paid': True}})},
            ]
            body = {'transaction_id': 'payment-' + 'b' * 32, 'items': items}
            first = self.post('/v1/transactions/payment', body)
            second = self.post('/v1/transactions/payment', body)
            self.assertEqual(first.status_code, 200)
            self.assertEqual(second.status_code, 200)
            self.assertFalse(first.get_json()['idempotent'])
            self.assertTrue(second.get_json()['idempotent'])
        finally:
            server.shutdown()
            server.server_close()
            thread.join(timeout=5)

    def test_payment_transaction_conflict_changes_nothing(self):
        schedule = {'lesson': {'paid': False}}
        payments = {'old': {'amount': 500}}
        self.post('/v1/json/write', {'path': 'schedule.json', 'data': schedule})
        self.post('/v1/json/write', {'path': 'payments.json', 'data': payments})
        self.post('/v1/files/write', {
            'path': 'book.xlsx', 'data': base64.b64encode(b'old-book').decode(),
            'expected_version': None,
        })
        response = self.post('/v1/transactions/payment', {
            'transaction_id': 'payment-' + 'c' * 32,
            'items': [
                {'path': 'schedule.json', 'kind': 'json', 'data': {'changed': True},
                 'expected_version': version_for(schedule)},
                {'path': 'payments.json', 'kind': 'json', 'data': {'changed': True},
                 'expected_version': '0' * 64},
                {'path': 'book.xlsx', 'kind': 'binary',
                 'data': base64.b64encode(b'new-book').decode(),
                 'expected_version': hashlib.sha256(b'old-book').hexdigest()},
            ],
        })
        self.assertEqual(response.status_code, 409)
        self.assertEqual(json.loads((self.root / 'schedule.json').read_text(encoding='utf-8')), schedule)
        self.assertEqual(json.loads((self.root / 'payments.json').read_text(encoding='utf-8')), payments)
        self.assertEqual((self.root / 'book.xlsx').read_bytes(), b'old-book')

    def test_payment_transaction_rejects_mixed_tenants_and_requires_auth(self):
        body = {
            'transaction_id': 'payment-' + 'f' * 32,
            'items': [
                {'path': 'teacher_data/101/schedule.json', 'kind': 'json',
                 'data': {'paid': True}, 'expected_version': None},
                {'path': 'teacher_data/202/book.xlsx', 'kind': 'binary',
                 'data': base64.b64encode(b'book').decode(), 'expected_version': None},
            ],
        }
        self.assertEqual(self.post('/v1/transactions/payment', body, {}).status_code, 401)
        self.assertEqual(self.post('/v1/transactions/payment', body).status_code, 400)
        self.assertFalse((self.root / 'teacher_data/101/schedule.json').exists())
        self.assertFalse((self.root / 'teacher_data/202/book.xlsx').exists())

    def test_payment_transaction_rolls_back_all_files_on_write_failure(self):
        schedule = {'lesson': {'paid': False}}
        payments = {'old': {'amount': 500}}
        self.post('/v1/json/write', {'path': 'schedule.json', 'data': schedule})
        self.post('/v1/json/write', {'path': 'payments.json', 'data': payments})
        self.post('/v1/files/write', {
            'path': 'book.xlsx', 'data': base64.b64encode(b'old-book').decode(),
            'expected_version': None,
        })
        import storage_server
        original_atomic_bytes = storage_server._atomic_bytes
        failed = {'value': False}

        def fail_first_book_write(path, raw):
            if Path(path).name == 'book.xlsx' and not failed['value']:
                failed['value'] = True
                raise OSError('simulated book failure')
            return original_atomic_bytes(path, raw)

        body = {
            'transaction_id': 'payment-' + 'd' * 32,
            'items': [
                {'path': 'schedule.json', 'kind': 'json', 'data': {'lesson': {'paid': True}},
                 'expected_version': version_for(schedule)},
                {'path': 'payments.json', 'kind': 'json', 'data': {'new': {'amount': 900}},
                 'expected_version': version_for(payments)},
                {'path': 'book.xlsx', 'kind': 'binary',
                 'data': base64.b64encode(b'new-book').decode(),
                 'expected_version': hashlib.sha256(b'old-book').hexdigest()},
            ],
        }
        with patch('storage_server._atomic_bytes', side_effect=fail_first_book_write):
            response = self.post('/v1/transactions/payment', body)
        self.assertEqual(response.status_code, 503)
        self.assertEqual(json.loads((self.root / 'schedule.json').read_text(encoding='utf-8')), schedule)
        self.assertEqual(json.loads((self.root / 'payments.json').read_text(encoding='utf-8')), payments)
        self.assertEqual((self.root / 'book.xlsx').read_bytes(), b'old-book')

    def test_prepared_payment_transaction_recovers_before_server_starts(self):
        after_schedule = {'lesson': {'paid': True}}
        (self.root / 'schedule.json').write_text(
            json.dumps(after_schedule), encoding='utf-8',
        )
        journal_dir = self.root / '.payment-transactions'
        journal_dir.mkdir()
        journal = journal_dir / ('payment-' + 'e' * 32 + '.json')
        journal.write_text(json.dumps({
            'schema': 1,
            'state': 'prepared',
            'transaction_id': 'payment-' + 'e' * 32,
            'request_hash': 'f' * 64,
            'items': [
                {'path': 'schedule.json', 'kind': 'json', 'existed': True,
                 'before': {}, 'after_version': version_for(after_schedule)},
                {'path': 'book.xlsx', 'kind': 'binary', 'existed': False,
                 'before': None, 'after_version': hashlib.sha256(b'new-book').hexdigest()},
            ],
            'versions': {},
        }), encoding='utf-8')
        create_app(self.root, self.token, backups=self.backups)
        self.assertEqual(json.loads((self.root / 'schedule.json').read_text(encoding='utf-8')), {})
        self.assertFalse((self.root / 'book.xlsx').exists())
        self.assertEqual(json.loads(journal.read_text(encoding='utf-8'))['state'], 'rolled_back')

    def test_binary_size_limit_and_tenant_separation(self):
        with patch('storage_server.MAX_FILE_BYTES', 4):
            response = self.post('/v1/files/write', {'path': 'book.xlsx', 'data': base64.b64encode(b'12345').decode(), 'expected_version': None})
            self.assertEqual(response.status_code, 400)
            self.assertFalse((self.root / 'book.xlsx').exists())
        for teacher, raw in [('101', b'one'), ('202', b'two')]:
            path = 'teacher_data/' + teacher + '/book.xlsx'
            self.assertEqual(self.post('/v1/files/write', {'path': path, 'data': base64.b64encode(raw).decode(), 'expected_version': None}).status_code, 200)
            self.assertEqual((self.root / path).read_bytes(), raw)
        self.assertEqual((self.root / 'teacher_data/101/book.xlsx').read_bytes(), b'one')

    def test_unsafe_or_tampered_binary_archive_cannot_restore(self):
        original, _ = create_backup(self.root, self.backups)
        for relative, content, checksum in [('../escape.pdf', b'x', hashlib.sha256(b'x').hexdigest()),
                                             ('receipts/check_1.pdf', b'x', '0' * 64)]:
            damaged = self.backups / 'damaged.zip'
            with zipfile.ZipFile(original) as source, zipfile.ZipFile(damaged, 'w') as target:
                for name in source.namelist():
                    raw = source.read(name)
                    if name == 'manifest.json':
                        manifest = json.loads(raw)
                        manifest['files'].append({'path': relative, 'size': 1, 'sha256': checksum})
                        raw = json.dumps(manifest).encode()
                    target.writestr(name, raw)
                target.writestr('data/' + relative, content)
            before = (self.root / 'students.json').read_bytes()
            with self.assertRaises(StorageServiceError):
                restore_backup(self.root, damaged, self.backups)
            self.assertEqual((self.root / 'students.json').read_bytes(), before)

    def test_client_rejects_file_checksum_mismatch(self):
        remote = RemoteJsonStorage('https://storage.example', self.token)
        with patch.object(remote, '_request', return_value={
            'status': 'ok', 'exists': True, 'data': 'eA==', 'version': '0' * 64,
        }):
            with self.assertRaisesRegex(RemoteStorageError, 'file_checksum_mismatch'):
                remote.read_file('book.xlsx')
        self.assertNotIn('book.xlsx', remote._file_versions)

    def test_rejects_unauthorized_and_path_traversal_requests(self):
        self.assertEqual(
            self.post("/v1/json/read", {"path": "schedule.json"}, {}).status_code,
            401,
        )
        for path in ("../secret.json", "/etc/passwd", "teacher_data/../../secret.json", "book.xlsx"):
            with self.subTest(path=path):
                self.assertEqual(
                    self.post("/v1/json/read", {"path": path}).status_code,
                    400,
                )

    def test_optimistic_version_prevents_lost_update(self):
        first = {"version": 1}
        self.assertEqual(
            self.post("/v1/json/write", {"path": "schedule.json", "data": first}).status_code,
            200,
        )
        stale = version_for(first)
        second = {"version": 2}
        self.assertEqual(
            self.post("/v1/json/write", {
                "path": "schedule.json",
                "data": second,
                "expected_version": stale,
            }).status_code,
            200,
        )
        conflict = self.post("/v1/json/write", {
            "path": "schedule.json",
            "data": {"version": 3},
            "expected_version": stale,
        })
        self.assertEqual(conflict.status_code, 409)
        self.assertEqual(conflict.get_json()["code"], "version_conflict")
        stored = json.loads((self.root / "schedule.json").read_text(encoding="utf-8"))
        self.assertEqual(stored, second)

    def test_refuses_implicit_empty_or_existing_unprepared_storage(self):
        other = Path(self.temp.name) / "other"
        with self.assertRaises(StorageServiceError):
            create_app(other, self.token)
        other.mkdir()
        (other / "unknown.txt").write_text("keep", encoding="utf-8")
        with self.assertRaises(StorageServiceError):
            create_app(other, self.token, initialize=True)

    def test_real_client_uses_token_and_detects_concurrent_write(self):
        server = make_server("127.0.0.1", 0, self.app)
        thread = threading.Thread(target=server.serve_forever, daemon=True)
        thread.start()
        try:
            url = "http://127.0.0.1:" + str(server.server_port)
            with patch.dict(os.environ, TEMLI_STORAGE_ALLOW_HTTP="true"):
                first = RemoteJsonStorage(url, self.token)
                second = RemoteJsonStorage(url, self.token)
                self.assertEqual(first.read_json("students.json", {}), {})
                self.assertEqual(second.read_json("students.json", {}), {})
                first.write_json("students.json", {"one": {"name": "Тест"}})
                with self.assertRaises(RemoteStorageConflict):
                    second.write_json("students.json", {"two": {"name": "Другой"}})
                self.assertEqual(
                    first.read_json("students.json", {}),
                    {"one": {"name": "Тест"}},
                )
        finally:
            server.shutdown()
            thread.join(timeout=5)

    def test_client_requires_https_and_complete_configuration(self):
        with self.assertRaises(RemoteStorageError):
            RemoteJsonStorage("http://storage.example", self.token)
        with self.assertRaises(RemoteStorageError):
            RemoteJsonStorage("https://storage.example", "short")

    def test_real_bot_process_reads_and_writes_only_remote_json(self):
        server = make_server("127.0.0.1", 0, self.app)
        thread = threading.Thread(target=server.serve_forever, daemon=True)
        thread.start()
        scratch = Path(self.temp.name) / "foreign-scratch"
        env = dict(os.environ)
        env.pop("DATA_DIR", None)
        env.pop("TEMLI_DATA_DIR", None)
        env.update({
            "TEMLI_STORAGE_URL": "http://127.0.0.1:" + str(server.server_port),
            "TEMLI_STORAGE_TOKEN": self.token,
            "TEMLI_STORAGE_ALLOW_HTTP": "true",
            "TEMLI_REMOTE_SCRATCH_DIR": str(scratch),
            "SCHEDULE_OWNER_ID": "101",
            "SCHEDULE_BOT_TOKEN": "",
        })
        code = """
import base64
import os
import bot
bot.ensure_teacher_registered('101')
with bot.teacher_scope('101'):
    bot.save_json(bot.STUDENTS_FILE, {'student-1': {'name': 'Remote only'}})
    assert bot.load_json(bot.STUDENTS_FILE)['student-1']['name'] == 'Remote only'
    image = base64.b64decode('iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk+A8AAQUBAScY42YAAAAASUVORK5CYII=')
    with bot.flask_app.app_context():
        response = bot.save_receipt_asset('logo', '.png', image, 'receipt_logo')
        assert response.get_json()['status'] == 'ok'
    settings = bot.load_settings()
    local_asset = bot.get_receipt_asset_path(settings, 'receipt_logo')
    assert open(local_asset, 'rb').read() == image
    replacement = b'jpeg-test-content'
    with bot.flask_app.app_context():
        response = bot.save_receipt_asset('logo', '.jpg', replacement, 'receipt_logo')
        assert response.get_json()['filename'] == 'logo.jpg'
    settings = bot.load_settings()
    local_asset = bot.get_receipt_asset_path(settings, 'receipt_logo')
    assert open(local_asset, 'rb').read() == replacement
    schedule = bot.load_json(bot.DATA_FILE, {})
    history = bot.load_json(bot.payments_file(), {})
    with bot.payment_files_transaction():
        bot.add_receipt_to_book('Remote client', 1250, 'test-001', bot.receipt_now())
        bot.add_receipt_to_book('Second row', 750, 'test-002', bot.receipt_now())
        schedule['lesson-1'] = {'paid': True}
        history['payment-1'] = {'amount': 1250}
        bot.save_json(bot.DATA_FILE, schedule)
        bot.save_json(bot.payments_file(), history)
    assert bot.load_json(bot.DATA_FILE)['lesson-1']['paid'] is True
    assert bot.load_json(bot.payments_file())['payment-1']['amount'] == 1250
bot.init_book()
assert not os.path.exists(bot.STUDENTS_FILE)
assert os.path.exists(bot.BOOK_FILE)
print('remote bot storage OK')
"""
        try:
            result = subprocess.run(
                [sys.executable, "-c", code],
                env=env,
                capture_output=True,
                timeout=90,
            )
        finally:
            server.shutdown()
            thread.join(timeout=5)
        self.assertEqual(result.returncode, 0, result.stderr.decode(errors="replace"))
        students = json.loads((self.root / "students.json").read_text(encoding="utf-8"))
        self.assertEqual(students["student-1"]["name"], "Remote only")
        self.assertFalse((scratch / "students.json").exists())
        self.assertFalse((self.root / 'receipt_assets/logo.png').exists())
        self.assertTrue((self.root / 'receipt_assets/logo.jpg').is_file())
        self.assertFalse((scratch / 'receipt_assets/logo.png').exists())
        self.assertTrue((scratch / 'receipt_assets/logo.jpg').is_file())
        workbook = openpyxl.load_workbook(
            io.BytesIO((self.root / 'book.xlsx').read_bytes()), data_only=True,
        )
        try:
            self.assertEqual(workbook.active.cell(row=2, column=4).value, 'Remote client')
            self.assertEqual(workbook.active.cell(row=2, column=5).value, 1250)
            self.assertEqual(workbook.active.cell(row=3, column=4).value, 'Second row')
            self.assertEqual(workbook.active.cell(row=3, column=5).value, 750)
        finally:
            workbook.close()

    def test_backup_api_creates_lists_and_downloads_verified_archive(self):
        self.assertEqual(
            self.post("/v1/json/write", {
                "path": "teacher_data/101/students.json",
                "data": {"student-1": {"name": "Тест"}},
            }).status_code,
            200,
        )
        created = self.post("/v1/backups", {}).get_json()
        self.assertEqual(created["status"], "ok")
        name = created["name"]

        listed = self.client.get("/v1/backups", headers=self.headers)
        self.assertEqual(listed.status_code, 200)
        item = listed.get_json()["backups"][0]
        self.assertEqual(item["name"], name)
        self.assertEqual(item["sha256"], created["sha256"])
        self.assertGreaterEqual(item["file_count"], 6)

        downloaded = self.client.get("/v1/backups/" + name, headers=self.headers)
        self.addCleanup(downloaded.close)
        self.assertEqual(downloaded.status_code, 200)
        with zipfile.ZipFile(io.BytesIO(downloaded.data)) as archive:
            manifest = json.loads(archive.read("manifest.json"))
            members = {entry["path"]: entry for entry in manifest["files"]}
            entry = members["teacher_data/101/students.json"]
            raw = archive.read("data/teacher_data/101/students.json")
            self.assertEqual(len(raw), entry["size"])
            self.assertEqual(json.loads(raw)["student-1"]["name"], "Тест")

        self.assertEqual(self.client.get("/v1/backups").status_code, 401)
        self.assertEqual(self.client.get("/v1/backups/../secret.zip").status_code, 401)

    def test_restore_requires_one_time_challenge_and_creates_safety_backup(self):
        original = {"student-1": {"name": "До изменения"}}
        changed = {"student-2": {"name": "После изменения"}}
        self.assertEqual(
            self.post("/v1/json/write", {"path": "students.json", "data": original}).status_code,
            200,
        )
        name = self.post("/v1/backups", {}).get_json()["name"]
        self.assertEqual(
            self.post("/v1/json/write", {"path": "students.json", "data": changed}).status_code,
            200,
        )

        refused = self.post("/v1/backups/" + name + "/restore", {"apply": True})
        self.assertEqual(refused.status_code, 409)
        self.assertEqual(
            json.loads((self.root / "students.json").read_text(encoding="utf-8")),
            changed,
        )

        challenge = self.post(
            "/v1/backups/" + name + "/restore-challenge", {}).get_json()
        self.assertGreaterEqual(len(challenge["nonce"]), 32)
        restored = self.post("/v1/backups/" + name + "/restore", {
            "apply": True, "nonce": challenge["nonce"],
        })
        self.assertEqual(restored.status_code, 200)
        self.assertTrue(restored.get_json()["safety_backup"].endswith("-pre-restore.zip"))
        self.assertEqual(
            json.loads((self.root / "students.json").read_text(encoding="utf-8")),
            original,
        )
        reasons = {item.get("reason") for item in self.client.get(
            "/v1/backups", headers=self.headers,
        ).get_json()["backups"]}
        self.assertIn("pre-restore", reasons)
        replay = self.post("/v1/backups/" + name + "/restore", {
            "apply": True, "nonce": challenge["nonce"],
        })
        self.assertEqual(replay.status_code, 409)

    def test_tampered_backup_is_rejected_without_changing_storage(self):
        self.backups.mkdir(parents=True, exist_ok=True)
        name = "temli-20260921T120000000000Z-manual.zip"
        path = self.backups / name
        manifest = {
            "schema": 1,
            "created_at": "2026-09-21T12:00:00+00:00",
            "reason": "manual",
            "files": [{"path": "students.json", "size": 2, "sha256": "0" * 64}],
        }
        with zipfile.ZipFile(path, "w") as archive:
            archive.writestr("manifest.json", json.dumps(manifest))
            archive.writestr("data/students.json", b"{}")
        before = (self.root / "students.json").read_bytes()
        response = self.post(
            "/v1/backups/" + name + "/restore-challenge", {})
        self.assertEqual(response.status_code, 400)
        self.assertEqual(response.get_json()["code"], "invalid_backup")
        self.assertEqual((self.root / "students.json").read_bytes(), before)

    def test_scoped_tokens_enforce_least_privilege(self):
        app_token = "app-" + "a" * 40
        backup_token = "backup-" + "b" * 40
        admin_token = "admin-" + "c" * 40
        scoped = create_app(
            Path(self.temp.name) / "scoped", initialize=True,
            backups=Path(self.temp.name) / "scoped-backups",
            app_token=app_token, backup_read_token=backup_token,
            admin_token=admin_token,
        ).test_client()
        headers = lambda token: {"Authorization": "Bearer " + token}
        created = scoped.post("/v1/backups", json={}, headers=headers(app_token))
        self.assertEqual(created.status_code, 200)
        name = created.get_json()["name"]

        self.assertEqual(scoped.get("/v1/backups", headers=headers(app_token)).status_code, 401)
        self.assertEqual(scoped.get("/v1/backups", headers=headers(backup_token)).status_code, 200)
        self.assertEqual(scoped.post(
            "/v1/json/read", json={"path": "students.json"},
            headers=headers(backup_token)).status_code, 401)
        self.assertEqual(scoped.post(
            "/v1/backups/" + name + "/restore-challenge", json={},
            headers=headers(app_token)).status_code, 401)
        self.assertEqual(scoped.post(
            "/v1/backups/" + name + "/restore-challenge", json={},
            headers=headers(backup_token)).status_code, 401)
        challenge = scoped.post(
            "/v1/backups/" + name + "/restore-challenge", json={},
            headers=headers(admin_token))
        self.assertEqual(challenge.status_code, 200)
        nonce = challenge.get_json()["nonce"]
        self.assertEqual(scoped.post(
            "/v1/backups/" + name + "/restore", json={"apply": True, "nonce": nonce},
            headers=headers(backup_token)).status_code, 401)
        self.assertEqual(scoped.post(
            "/v1/backups/" + name + "/restore", json={"apply": True, "nonce": nonce},
            headers=headers(admin_token)).status_code, 200)

    def test_scoped_environment_resolution_and_legacy_fallback(self):
        legacy = "legacy-" + "x" * 40
        self.assertEqual(set(configured_storage_tokens({"TEMLI_STORAGE_TOKEN": legacy}).values()),
                         {legacy})
        values = configured_storage_tokens({
            "TEMLI_STORAGE_APP_TOKEN": "a" * 32,
            "TEMLI_STORAGE_BACKUP_READ_TOKEN": "b" * 32,
            "TEMLI_STORAGE_ADMIN_TOKEN": "c" * 32,
        })
        self.assertEqual(values, {"app": "a" * 32, "backup_read": "b" * 32,
                                  "admin": "c" * 32})
        configured = configured_remote_storage({
            "TEMLI_STORAGE_URL": "https://storage.example",
            "TEMLI_STORAGE_APP_TOKEN": "a" * 32,
            "TEMLI_STORAGE_BACKUP_READ_TOKEN": "b" * 32,
        })
        self.assertEqual(configured.token, "a" * 32)
        self.assertEqual(configured.backup_read_token, "b" * 32)
        with self.assertRaisesRegex(RemoteStorageError, "BACKUP_READ_TOKEN"):
            configured_remote_storage({
                "TEMLI_STORAGE_URL": "https://storage.example",
                "TEMLI_STORAGE_APP_TOKEN": "a" * 32,
            })

    def test_backup_rotation_keeps_only_newest_archives(self):
        for reason in ("one", "two", "three"):
            create_backup(self.root, self.backups, reason=reason, retention=2)
        archives = sorted(self.backups.glob("temli-*.zip"))
        self.assertEqual(len(archives), 2)
        self.assertTrue(any(path.name.endswith("-three.zip") for path in archives))


if __name__ == "__main__":
    unittest.main()
