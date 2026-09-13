"""Verified encrypted backups. No bot import or network during restore."""
import argparse
import datetime
import hashlib
import json
import os
import re
import tarfile
import tempfile
import threading
import time
import urllib.error
import urllib.parse
import urllib.request
import uuid
from contextlib import contextmanager
from pathlib import Path, PurePosixPath
from cryptography.fernet import Fernet
from verify_live_storage import inspect_storage

REQUIRED_ENV = ("TEMLI_GOOGLE_CLIENT_ID", "TEMLI_GOOGLE_CLIENT_SECRET",
                "TEMLI_GOOGLE_REFRESH_TOKEN", "TEMLI_GOOGLE_FOLDER_ID",
                "TEMLI_BACKUP_ENCRYPTION_KEY")
WEBHOOK_ENV = ("TEMLI_BACKUP_WEBHOOK_URL", "TEMLI_BACKUP_WEBHOOK_SECRET",
               "TEMLI_BACKUP_ENCRYPTION_KEY")
PREFIX = "temli-backup-"
NAME = re.compile(r"temli-backup-\d{8}-\d{6}(?:-[a-f0-9]{8})?\.tar\.gz\.enc")
FIELDS = "id,name,size,md5Checksum,parents,appProperties,trashed"
_THREAD_LOCK = threading.Lock()


def configured(environ=None):
    values = os.environ if environ is None else environ
    oauth = all(values.get(name, "").strip() for name in REQUIRED_ENV)
    webhook = all(values.get(name, "").strip() for name in WEBHOOK_ENV)
    return oauth or webhook


def webhook_configured(environ=None):
    values = os.environ if environ is None else environ
    return all(values.get(name, "").strip() for name in WEBHOOK_ENV)


def _env(name):
    value = os.getenv(name, "").strip()
    if not value:
        raise RuntimeError("missing_" + name.lower())
    return value


def atomic_write(path, data):
    path = Path(path)
    fd, temporary = tempfile.mkstemp(prefix=".writing-", dir=path.parent)
    try:
        with os.fdopen(fd, "wb") as target:
            target.write(data)
            target.flush()
            os.fsync(target.fileno())
        os.replace(temporary, path)
    finally:
        Path(temporary).unlink(missing_ok=True)


def read_json(path):
    try:
        value = json.loads(Path(path).read_text(encoding="utf-8"))
        return value if isinstance(value, dict) else {}
    except (OSError, ValueError):
        return {}


@contextmanager
def backup_lock(directory):
    """Serialize worker and CLI, also across processes; never locks live data over network."""
    with _THREAD_LOCK:
        with open(Path(directory) / ".backup.lock", "a+b") as handle:
            if os.name == "nt":
                import msvcrt
                if handle.seek(0, 2) == 0:
                    handle.write(b"0")
                    handle.flush()
                handle.seek(0)
                msvcrt.locking(handle.fileno(), msvcrt.LK_NBLCK, 1)
            else:
                import fcntl
                fcntl.flock(handle.fileno(), fcntl.LOCK_EX | fcntl.LOCK_NB)
            try:
                yield
            finally:
                if os.name == "nt":
                    handle.seek(0)
                    msvcrt.locking(handle.fileno(), msvcrt.LK_UNLCK, 1)
                else:
                    fcntl.flock(handle.fileno(), fcntl.LOCK_UN)


def _token():
    data = urllib.parse.urlencode({
        "client_id": _env("TEMLI_GOOGLE_CLIENT_ID"),
        "client_secret": _env("TEMLI_GOOGLE_CLIENT_SECRET"),
        "refresh_token": _env("TEMLI_GOOGLE_REFRESH_TOKEN"),
        "grant_type": "refresh_token",
    }).encode()
    try:
        with urllib.request.urlopen(urllib.request.Request(
                "https://oauth2.googleapis.com/token", data=data), timeout=30) as response:
            return json.load(response)["access_token"]
    except urllib.error.HTTPError as error:
        # Do not expose the response body or credentials.
        raise RuntimeError("google_authorization_failed_http_" + str(error.code)) from None


def _request(url, token, data=None, method=None, content_type=None):
    headers = {"Authorization": "Bearer " + token}
    if content_type:
        headers["Content-Type"] = content_type
    with urllib.request.urlopen(urllib.request.Request(
            url, data=data, headers=headers, method=method), timeout=120) as response:
        raw = response.read()
        return json.loads(raw) if raw else {}


def extract_checked(path, destination):
    with tarfile.open(path, "r:gz") as archive:
        members = archive.getmembers()
        seen = set()
        size = 0
        for member in members:
            part = PurePosixPath(member.name)
            folded = str(part).casefold()
            size += member.size
            if (part.is_absolute() or ".." in part.parts or "\\" in member.name
                    or ":" in member.name or not part.parts or part.parts[0] != "temli"
                    or not (member.isfile() or member.isdir()) or folded in seen
                    or size > 2 * 1024**3 or len(members) > 100000):
                raise RuntimeError("unsafe_archive_entry_or_size")
            seen.add(folded)
        archive.extractall(destination, filter="data")


def check_restored(root, inspector=inspect_storage):
    report = inspector(root, check_key=False)
    if report.get("status") != "ok":
        raise RuntimeError("backup_verification_failed:" + ",".join(report.get("errors", [])))
    from openpyxl import load_workbook
    for path in Path(root).rglob("*.xlsx"):
        book = load_workbook(path, read_only=True)
        try:
            for sheet in book:
                for row in sheet.iter_rows(values_only=True):
                    pass
        finally:
            book.close()
    report["personal_bot_decryption_checked"] = False
    return report


def verify_plain_archive(path, inspector=inspect_storage):
    with tempfile.TemporaryDirectory(prefix="temli-backup-check-") as temporary:
        extract_checked(path, temporary)
        return check_restored(Path(temporary) / "temli", inspector)


def remote_matches(meta, path, content):
    return (meta.get("name") == path.name
            and str(meta.get("size")) == str(len(content))
            and meta.get("md5Checksum") == hashlib.md5(content).hexdigest()
            and _env("TEMLI_GOOGLE_FOLDER_ID") in meta.get("parents", [])
            and not meta.get("trashed", False))


def remote_files(token):
    folder = _env("TEMLI_GOOGLE_FOLDER_ID")
    if not re.fullmatch(r"[A-Za-z0-9_-]+", folder):
        raise RuntimeError("invalid_folder_id")
    query = "'" + folder + "' in parents and trashed=false"
    page = None
    result = []
    while True:
        params = {"q": query, "fields": "nextPageToken,files(" + FIELDS + ")",
                  "pageSize": 1000}
        if page:
            params["pageToken"] = page
        data = _request("https://www.googleapis.com/drive/v3/files?" +
                        urllib.parse.urlencode(params), token)
        result.extend(data.get("files", []))
        page = data.get("nextPageToken")
        if not page:
            return result


def _upload_webhook(path):
    url = _env("TEMLI_BACKUP_WEBHOOK_URL")
    if not url.startswith("https://script.google.com/macros/s/") or not url.endswith("/exec"):
        raise RuntimeError("invalid_backup_webhook_url")
    content = path.read_bytes()
    checksum = hashlib.sha256(content).hexdigest()
    payload = json.dumps({
        "secret": _env("TEMLI_BACKUP_WEBHOOK_SECRET"), "name": path.name,
        "sha256": checksum,
        "content_base64": __import__("base64").b64encode(content).decode("ascii"),
    }).encode()
    try:
        request = urllib.request.Request(url, data=payload, headers={
            "Content-Type": "application/json", "User-Agent": "TEMLI-backup/1",
        })
        with urllib.request.urlopen(request, timeout=120) as response:
            result = json.loads(response.read())
    except (OSError, ValueError, urllib.error.HTTPError) as error:
        raise RuntimeError("backup_webhook_delivery_failed_" + type(error).__name__) from None
    if (result.get("status") != "ok" or result.get("name") != path.name
            or str(result.get("size")) != str(len(content))
            or result.get("sha256") != checksum
            or not isinstance(result.get("id"), str)):
        raise RuntimeError("backup_webhook_verification_failed")
    return result


def _upload(path):
    if webhook_configured():
        return _upload_webhook(path)
    token = _token()
    content = path.read_bytes()
    # A previous upload may have completed before the connection failed.
    for candidate in remote_files(token):
        if remote_matches(candidate, path, content):
            return candidate
    boundary = "temli-" + uuid.uuid4().hex
    metadata = json.dumps({
        "name": path.name, "parents": [_env("TEMLI_GOOGLE_FOLDER_ID")],
        "appProperties": {"temli_backup": "v2"},
    }).encode()
    marker = boundary.encode()
    body = (b"--" + marker + b"\r\nContent-Type: application/json\r\n\r\n" + metadata
            + b"\r\n--" + marker + b"\r\nContent-Type: application/octet-stream\r\n\r\n"
            + content + b"\r\n--" + marker + b"--\r\n")
    created = _request(
        "https://www.googleapis.com/upload/drive/v3/files?uploadType=multipart&fields=id",
        token, body, "POST", "multipart/related; boundary=" + boundary)
    meta = _request("https://www.googleapis.com/drive/v3/files/" + created["id"]
                    + "?fields=" + FIELDS, token)
    if not remote_matches(meta, path, content):
        raise RuntimeError("google_uploaded_checksum_or_size_mismatch")
    return meta


def rotate_remote(current_id):
    token = _token()
    # Only this module's marked archives; legacy or unrelated files are preserved.
    files = [f for f in remote_files(token) if NAME.fullmatch(f.get("name", ""))
             and f.get("appProperties", {}).get("temli_backup") == "v2"]
    files.sort(key=lambda f: f["name"], reverse=True)
    keep = {current_id}
    for f in files:
        if len(keep) >= 30:
            break
        keep.add(f["id"])
    for f in files:
        if f["id"] not in keep:
            _request("https://www.googleapis.com/drive/v3/files/" + f["id"], token,
                     b'{"trashed":true}', "PATCH", "application/json")


def local_directory(host):
    root = Path(host.BASE_DIR).resolve()
    output = Path(os.getenv("TEMLI_BACKUP_LOCAL_DIR", str(root.parent / "temli-backups"))).resolve()
    if output == root or output.is_relative_to(root) or root.is_relative_to(output):
        raise RuntimeError("backup_directory_overlaps_storage")
    output.mkdir(parents=True, exist_ok=True)
    return output


def create(host, now=None, upload=True, force=False):
    now = now or datetime.datetime.now(__import__("pytz").timezone(
        getattr(host, "TIMEZONE_NAME", "Europe/Moscow")))
    output = local_directory(host)
    cipher = Fernet(_env("TEMLI_BACKUP_ENCRYPTION_KEY"))
    with backup_lock(output):
        state_path = output / "last-success.json"
        state = read_json(state_path)
        if (upload and not force and state.get("date") == now.date().isoformat()
                and state.get("drive_file_id")):
            return dict(state, status="already_done")
        pending_path = output / "pending.json"
        pending = read_json(pending_path)
        final = output / pending.get("archive", "invalid")
        valid_pending = (pending.get("archive") == final.name and NAME.fullmatch(final.name)
                         and final.is_file() and not final.is_symlink()
                         and pending.get("sha256") == hashlib.sha256(final.read_bytes()).hexdigest())
        if pending and not valid_pending:
            raise RuntimeError("pending_archive_invalid")
        if valid_pending:
            cipher.decrypt(final.read_bytes())  # reject a changed key before marking success
            report = pending["counts"]
        else:
            stamp = now.strftime("%Y%m%d-%H%M%S") + "-" + uuid.uuid4().hex[:8]
            final = output / (PREFIX + stamp + ".tar.gz.enc")
            with tempfile.TemporaryDirectory(prefix=".snapshot-", dir=output) as temporary:
                plain = Path(temporary) / "snapshot.tar.gz"
                with host.DATA_LOCK:
                    with tarfile.open(plain, "w:gz") as archive:
                        archive.add(host.BASE_DIR, arcname="temli")
                report = verify_plain_archive(plain)
                data = plain.read_bytes()
                encrypted = cipher.encrypt(data)
                if cipher.decrypt(encrypted) != data:
                    raise RuntimeError("encryption_roundtrip_failed")
                atomic_write(final, encrypted)
            pending = {"archive": final.name, "date": now.date().isoformat(),
                       "sha256": hashlib.sha256(final.read_bytes()).hexdigest(),
                       "counts": {k: report[k] for k in ("lesson_records", "student_records")}}
            atomic_write(pending_path, json.dumps(pending).encode())
        if not upload:
            return dict(status="local_only", **pending)
        remote = _upload(final)
        state = {"date": pending["date"], "archive": final.name, "sha256": pending["sha256"],
                 "drive_file_id": remote["id"], **pending["counts"]}
        atomic_write(state_path, json.dumps(state, indent=2).encode())
        pending_path.unlink(missing_ok=True)
        warnings = []
        try:
            if not webhook_configured():
                rotate_remote(remote["id"])
            files = sorted((f for f in output.iterdir() if NAME.fullmatch(f.name)
                            and f.is_file() and not f.is_symlink()), key=lambda f: f.name, reverse=True)
            keep = {final}
            for f in files:
                if len(keep) >= 7:
                    break
                keep.add(f)
            for f in files:
                if f not in keep:
                    f.unlink()
        except Exception as error:
            warnings.append("rotation_failed_" + type(error).__name__)
        return dict(state, status="ok", warnings=warnings)


def restore(encrypted, destination, confirm):
    if confirm != "RESTORE":
        raise RuntimeError("confirmation_required")
    raw_destination = Path(destination).absolute()
    if raw_destination.exists() or raw_destination.is_symlink():
        raise RuntimeError("destination_must_not_exist")
    destination = raw_destination.resolve()
    destination.parent.mkdir(parents=True, exist_ok=True)
    # Same-filesystem staging; target only appears after full verification.
    with tempfile.TemporaryDirectory(prefix=".temli-restore-", dir=destination.parent) as temporary:
        temporary = Path(temporary)
        plain = temporary / "snapshot.tar.gz"
        plain.write_bytes(Fernet(_env("TEMLI_BACKUP_ENCRYPTION_KEY")).decrypt(Path(encrypted).read_bytes()))
        stage = temporary / "stage"
        stage.mkdir()
        extract_checked(plain, stage)
        report = check_restored(stage / "temli")
        if destination.exists():
            raise RuntimeError("destination_must_not_exist")
        os.rename(stage, destination)
    report["storage"] = str(destination / "temli")
    return report


def status(host):
    output = local_directory(host)
    return {"configured": configured(), "last_success": read_json(output / "last-success.json"),
            "pending_upload": bool(read_json(output / "pending.json")),
            "last_attempt": read_json(output / "last-attempt.json")}


def start_worker(host):
    if not configured():
        print("TEMLI automatic backup: disabled (environment is incomplete)", flush=True)
        return None
    def work():
        while True:
            try:
                result = create(host)
                if result["status"] != "already_done":
                    print("TEMLI automatic backup: " + result["status"] + " " + result["archive"], flush=True)
                atomic_write(local_directory(host) / "last-attempt.json", json.dumps({
                    "at": datetime.datetime.now(datetime.timezone.utc).isoformat(),
                    "status": result["status"], "warnings": result.get("warnings", [])}).encode())
            except Exception as error:
                print("TEMLI automatic backup: failed " + type(error).__name__, flush=True)
                try:
                    atomic_write(local_directory(host) / "last-attempt.json", json.dumps({
                        "at": datetime.datetime.now(datetime.timezone.utc).isoformat(),
                        "status": "failed", "error_type": type(error).__name__}).encode())
                except Exception:
                    pass
            time.sleep(900)
    thread = threading.Thread(target=work, name="temli-backup-worker", daemon=True)
    thread.start()
    return thread


def main():
    parser = argparse.ArgumentParser()
    sub = parser.add_subparsers(dest="command", required=True)
    create_cmd = sub.add_parser("create")
    create_cmd.add_argument("--force", action="store_true")
    sub.add_parser("status")
    cmd = sub.add_parser("restore")
    cmd.add_argument("archive")
    cmd.add_argument("--destination", required=True)
    cmd.add_argument("--confirm", required=True)
    args = parser.parse_args()
    try:
        if args.command == "restore":
            result = restore(args.archive, args.destination, args.confirm)
        else:
            import bot
            result = create(bot, force=args.force) if args.command == "create" else status(bot)
        print(json.dumps(result, ensure_ascii=True, indent=2))
    except Exception as error:
        print(json.dumps({"status": "error", "error_type": type(error).__name__}))
        raise SystemExit(1)

if __name__ == "__main__":
    main()
