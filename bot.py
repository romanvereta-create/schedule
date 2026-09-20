import os
import sys
import subprocess

DEPENDENCIES = [
    ("Flask", "flask"),
    ("python-telegram-bot", "telegram"),
    ("pytz", "pytz"),
    ("Flask-Cors", "flask_cors"),
    ("fpdf==1.7.2", "fpdf"),
    ("openpyxl==3.1.2", "openpyxl"),
    ("cryptography>=44,<47", "cryptography"),
]

for package_name, import_name in DEPENDENCIES:
    try:
        __import__(import_name)
    except ImportError:
        try:
            subprocess.check_call([
                sys.executable, "-m", "pip", "install", package_name, "--quiet"
            ], timeout=120 if import_name == "cryptography" else None)
        except (subprocess.CalledProcessError, subprocess.TimeoutExpired):
            if import_name != "cryptography":
                raise
            print("Не удалось установить cryptography: подключение личных ботов временно недоступно.")

import asyncio
import copy
import contextvars
import datetime
import hashlib
import io
import hmac
import json
import math
import re
import tempfile
import threading
import time
import shutil
from contextlib import contextmanager
from functools import wraps
from urllib.parse import parse_qsl

if os.name == "nt":
    import msvcrt
else:
    import fcntl

import pytz
from flask import Flask, jsonify, request, send_file, g
from flask_cors import CORS
from telegram import InlineKeyboardButton, InlineKeyboardMarkup, Update, WebAppInfo
from telegram.ext import Application, CommandHandler, ContextTypes
from fpdf import FPDF
import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side

from persistent_storage import resolve_storage_root
from remote_storage import RemoteStorageError, configured_remote_storage
from calendar_undo import CalendarUndo

CALENDAR_UNDO = CalendarUndo()

CODE_DIR = os.path.dirname(os.path.abspath(__file__))
# BASE_DIR remains the data root for tenant helpers and personal-bot modules.
REMOTE_STORAGE = configured_remote_storage()
if REMOTE_STORAGE:
    # Only locks and temporary binary artifacts live here. JSON data is never
    # written to this directory while remote storage is enabled.
    BASE_DIR = os.path.abspath(os.getenv(
        "TEMLI_REMOTE_SCRATCH_DIR",
        os.path.join(tempfile.gettempdir(), "temli-remote"),
    ))
    os.makedirs(BASE_DIR, exist_ok=True)
else:
    BASE_DIR = resolve_storage_root(CODE_DIR)


def project_path(*parts):
    return os.path.join(BASE_DIR, *parts)


TOKEN = os.getenv("SCHEDULE_BOT_TOKEN")
DATA_FILE = project_path("schedule.json")
STUDENTS_FILE = project_path("students.json")
SETTINGS_FILE = project_path("settings.json")
WEBAPP_URL = "https://romanvereta-create.github.io/schedule-mini-app/"
WEBAPP_ORIGIN = os.getenv("SCHEDULE_WEBAPP_ORIGIN", "https://romanvereta-create.github.io")
OWNER_ID = os.getenv("SCHEDULE_OWNER_ID", "").strip()
TIMEZONE_NAME = os.getenv("SCHEDULE_TIMEZONE", "Europe/Moscow")
ALLOW_UNAUTHENTICATED = os.getenv("ALLOW_UNAUTHENTICATED", "false").lower() == "true"
RECEIPTS_DIR = project_path("receipts")
RECEIPT_ASSETS_DIR = project_path("receipt_assets")
BOOK_FILE = project_path("book.xlsx")
FONT_REGULAR = os.path.join(CODE_DIR, "DejaVuSansCondensed.ttf")
FONT_BOLD = os.path.join(CODE_DIR, "DejaVuSansCondensed-Bold.ttf")
TENANT_DATA_DIR = project_path("teacher_data")
TENANT_REGISTRY_FILE = project_path("teacher_registry.json")
TEACHER_CONTEXT = contextvars.ContextVar("schedule_teacher_id", default="")

BOT_APPLICATION = None
BOT_LOOP = None
REMINDER_TASK = None

class InterProcessRLock:
    """Re-entrant thread lock backed by a process-wide filesystem lock."""

    def __init__(self, lock_path):
        self._thread_lock = threading.RLock()
        self._local = threading.local()
        self._lock_path = lock_path

    def __enter__(self):
        self._thread_lock.acquire()
        depth = getattr(self._local, "depth", 0)
        if depth == 0:
            os.makedirs(os.path.dirname(self._lock_path), exist_ok=True)
            handle = open(self._lock_path, "a+b")
            try:
                if os.name == "nt":
                    handle.seek(0)
                    if handle.tell() == 0 and os.path.getsize(self._lock_path) == 0:
                        handle.write(b"0")
                        handle.flush()
                    handle.seek(0)
                    msvcrt.locking(handle.fileno(), msvcrt.LK_LOCK, 1)
                else:
                    fcntl.flock(handle.fileno(), fcntl.LOCK_EX)
            except Exception:
                handle.close()
                self._thread_lock.release()
                raise
            self._local.handle = handle
        self._local.depth = depth + 1
        return self

    def __exit__(self, exc_type, exc, traceback):
        depth = getattr(self._local, "depth", 1) - 1
        self._local.depth = depth
        if depth == 0:
            handle = self._local.handle
            try:
                if os.name == "nt":
                    handle.seek(0)
                    msvcrt.locking(handle.fileno(), msvcrt.LK_UNLCK, 1)
                else:
                    fcntl.flock(handle.fileno(), fcntl.LOCK_UN)
            finally:
                handle.close()
                del self._local.handle
        self._thread_lock.release()


class DataCorruptionError(RuntimeError):
    pass


DATA_LOCK = InterProcessRLock(project_path(".schedule_data.lock"))
HEX_COLOR_RE = re.compile(r"^#[0-9a-fA-F]{6}$")


def serialized_data(function):
    """Hold the shared lock across an entire read/modify/write operation."""
    @wraps(function)
    def wrapped(*args, **kwargs):
        with DATA_LOCK:
            return function(*args, **kwargs)
    return wrapped


def _load_json_raw(filename, default=None):
    if default is None:
        default = {}
    with DATA_LOCK:
        if REMOTE_STORAGE:
            try:
                relative = os.path.relpath(os.path.abspath(filename), BASE_DIR)
                return REMOTE_STORAGE.read_json(relative, default)
            except RemoteStorageError as exc:
                raise DataCorruptionError(
                    f"Не удалось прочитать удалённое хранилище ({exc})."
                ) from exc
        if os.path.exists(filename):
            try:
                with open(filename, "r", encoding="utf-8") as f:
                    return json.load(f)
            except (OSError, json.JSONDecodeError) as exc:
                raise DataCorruptionError(
                    f"Не удалось прочитать файл данных {os.path.basename(filename)}. "
                    "Запись остановлена; восстановите файл из резервной копии .bak."
                ) from exc
    return default


def _save_json_raw(filename, data):
    if REMOTE_STORAGE:
        with DATA_LOCK:
            try:
                relative = os.path.relpath(os.path.abspath(filename), BASE_DIR)
                REMOTE_STORAGE.write_json(relative, data)
                return
            except RemoteStorageError as exc:
                raise DataCorruptionError(
                    f"Не удалось записать удалённое хранилище ({exc})."
                ) from exc
    directory = os.path.dirname(os.path.abspath(filename)) or "."
    os.makedirs(directory, exist_ok=True)
    with DATA_LOCK:
        if os.path.exists(filename):
            try:
                with open(filename, "r", encoding="utf-8") as source:
                    json.load(source)
            except (OSError, json.JSONDecodeError) as exc:
                raise DataCorruptionError(
                    f"Файл данных {os.path.basename(filename)} повреждён; запись отменена."
                ) from exc
            backup_fd, backup_temp = tempfile.mkstemp(prefix="schedule_backup_", suffix=".tmp", dir=directory)
            try:
                with os.fdopen(backup_fd, "wb") as backup_target, open(filename, "rb") as source:
                    shutil.copyfileobj(source, backup_target)
                    backup_target.flush()
                    os.fsync(backup_target.fileno())
                os.replace(backup_temp, f"{filename}.bak")
            finally:
                if os.path.exists(backup_temp):
                    os.remove(backup_temp)
        fd, temp_path = tempfile.mkstemp(prefix="schedule_", suffix=".tmp", dir=directory)
        try:
            with os.fdopen(fd, "w", encoding="utf-8") as f:
                json.dump(data, f, ensure_ascii=False, indent=2)
                f.flush()
                os.fsync(f.fileno())
            os.replace(temp_path, filename)
        finally:
            if os.path.exists(temp_path):
                os.remove(temp_path)


def _safe_teacher_id(value):
    value = str(value or "").strip()
    return re.sub(r"[^0-9A-Za-z_-]", "_", value) if value else ""


def legacy_root_data_exists():
    """Return True when assigning the root tenant implicitly could expose data."""
    if REMOTE_STORAGE:
        return any(bool(_load_json_raw(filename, {}))
                   for filename in (DATA_FILE, STUDENTS_FILE, SETTINGS_FILE))
    for filename in (DATA_FILE, STUDENTS_FILE, SETTINGS_FILE):
        if not os.path.exists(filename):
            continue
        try:
            with open(filename, "r", encoding="utf-8") as source:
                if bool(json.load(source)):
                    return True
        except (OSError, json.JSONDecodeError):
            return True
    return os.path.exists(BOOK_FILE) and os.path.getsize(BOOK_FILE) > 0


def load_tenant_registry():
    raw = _load_json_raw(TENANT_REGISTRY_FILE, {})
    if not isinstance(raw, dict):
        raw = {}
    teachers = raw.get("teachers") if isinstance(raw.get("teachers"), dict) else {}
    return {
        "primary_teacher_id": str(raw.get("primary_teacher_id", "") or "").strip(),
        "teachers": teachers,
    }


def save_tenant_registry(registry):
    _save_json_raw(TENANT_REGISTRY_FILE, registry)


def ensure_teacher_registered(teacher_id, user=None):
    teacher_id = str(teacher_id or "").strip()
    if not teacher_id:
        return ""

    with DATA_LOCK:
        registry = load_tenant_registry()
        primary = registry.get("primary_teacher_id", "")
        had_primary = bool(primary)
        if not primary:
            # Если SCHEDULE_OWNER_ID задан, старые корневые данные закрепляются за ним.
            # При существующих legacy-данных запрещаем назначать владельцем случайного
            # первого пользователя: оператор должен явно задать SCHEDULE_OWNER_ID.
            if not OWNER_ID and legacy_root_data_exists():
                raise RuntimeError(
                    "Найдены корневые legacy-данные, но SCHEDULE_OWNER_ID не задан. "
                    "Регистрация остановлена, чтобы не передать данные неверному преподавателю."
                )
            primary = OWNER_ID or teacher_id
            registry["primary_teacher_id"] = primary

        teachers = registry.setdefault("teachers", {})
        info = teachers.get(teacher_id, {}) if isinstance(teachers.get(teacher_id), dict) else {}
        previous_info = dict(info)
        if isinstance(user, dict):
            info.update({
                "id": teacher_id,
                "first_name": str(user.get("first_name", "") or ""),
                "last_name": str(user.get("last_name", "") or ""),
                "username": str(user.get("username", "") or ""),
            })
        else:
            info.setdefault("id", teacher_id)
        now = datetime.datetime.utcnow().replace(microsecond=0)
        try:
            seen = datetime.datetime.fromisoformat(str(previous_info.get('last_seen_at', '')).removesuffix('Z'))
            recently_seen = 0 <= (now - seen).total_seconds() < 60
        except (ValueError, TypeError):
            recently_seen = False
        # Opening the calendar issues several API reads. Do not fsync the global
        # registry on each one when neither identity nor minute-level activity changed.
        unchanged = had_primary and info == previous_info and recently_seen
        if not unchanged:
            info['last_seen_at'] = now.isoformat() + 'Z'
        teachers[teacher_id] = info
        if not unchanged:
            save_tenant_registry(registry)

        if teacher_id != primary:
            os.makedirs(os.path.join(TENANT_DATA_DIR, _safe_teacher_id(teacher_id)), exist_ok=True)
    return teacher_id


def registered_teacher_ids(include_legacy=True):
    registry = load_tenant_registry()
    ids = [str(x) for x in registry.get("teachers", {}).keys() if str(x).strip()]
    primary = str(registry.get("primary_teacher_id", "") or OWNER_ID or "").strip()
    if primary and primary not in ids:
        ids.insert(0, primary)
    if not ids and include_legacy:
        return [""]
    return ids


def current_teacher_id():
    return str(TEACHER_CONTEXT.get() or "").strip()


@contextmanager
def teacher_scope(teacher_id):
    token = TEACHER_CONTEXT.set(str(teacher_id or "").strip())
    try:
        yield
    finally:
        TEACHER_CONTEXT.reset(token)


def tenant_root(teacher_id=None):
    teacher_id = str(teacher_id if teacher_id is not None else current_teacher_id()).strip()
    registry = load_tenant_registry()
    primary = str(registry.get("primary_teacher_id", "") or OWNER_ID or "").strip()
    # До первой регистрации и для основного преподавателя сохраняем прежние корневые файлы.
    if not teacher_id or not primary or teacher_id == primary:
        return BASE_DIR
    return os.path.join(TENANT_DATA_DIR, _safe_teacher_id(teacher_id))


def tenant_file(filename, teacher_id=None):
    root = tenant_root(teacher_id)
    return filename if root == BASE_DIR else os.path.join(root, os.path.basename(filename))


def current_book_file():
    return tenant_file(BOOK_FILE)


def current_receipts_dir():
    return tenant_file(RECEIPTS_DIR)


def current_receipt_assets_dir():
    return tenant_file(RECEIPT_ASSETS_DIR)


def load_json(filename, default=None):
    path = tenant_file(filename) if filename in {DATA_FILE, STUDENTS_FILE, SETTINGS_FILE} else filename
    return _load_json_raw(path, default)


def save_json(filename, data):
    path = tenant_file(filename) if filename in {DATA_FILE, STUDENTS_FILE, SETTINGS_FILE} else filename
    _save_json_raw(path, data)

def normalize_contacts(data):
    contacts = data.get("contacts")
    if isinstance(contacts, dict):
        return {
            str(key): str(value).strip()
            for key, value in contacts.items()
            if key in {"tg", "wa", "phone", "max"} and str(value).strip()
        }

    legacy_value = str(data.get("parent_contact", "")).strip()
    legacy_type = str(data.get("parent_contact_type", "tg")).strip() or "tg"
    if legacy_value and legacy_type in {"tg", "wa", "phone", "max"}:
        return {legacy_type: legacy_value}
    return {}


def normalize_student_contacts(data):
    contacts = data.get("student_contacts")
    if not isinstance(contacts, dict):
        return {}
    return {
        str(key): str(value).strip()
        for key, value in contacts.items()
        if key in {"tg", "wa", "phone", "max"} and str(value).strip()
    }




def new_student_notification_settings():
    settings = load_settings()
    return {
        "student_reminders": settings.get("default_student_reminders", False) is True,
        "parent_lesson_end": settings.get("parent_lesson_end", False) is True,
    }


def normalize_group_members(data, students):
    raw_members = data.get("group_members")
    if not isinstance(raw_members, list):
        return []

    members = []
    seen = set()
    for raw in raw_members:
        if not isinstance(raw, dict):
            continue
        student_id = str(raw.get("student_id", "")).strip()
        name = str(raw.get("name", "")).strip()
        if student_id == "manual" or not student_id:
            if not name:
                continue
            existing_id = None
            for s_id, s_info in students.items():
                existing_name = s_info.get("name") if isinstance(s_info, dict) else s_info
                if existing_name == name:
                    existing_id = str(s_id)
                    break
            if existing_id:
                student_id = existing_id
            else:
                student_id = f"manual_{time.time_ns()}"
                students[student_id] = {
                    **new_student_notification_settings(),
                    "name": name,
                    "username": "",
                    "contacts": {},
                    "student_contacts": {},
                    "user_id": student_id,
                    "color": next_auto_color(students),
                }
        elif student_id in students:
            existing = students[student_id]
            existing_name = existing.get("name") if isinstance(existing, dict) else existing
            name = name or str(existing_name or student_id)
        elif name:
            students[student_id] = {
                **new_student_notification_settings(),
                "name": name,
                "username": "",
                "contacts": {},
                "student_contacts": {},
                "user_id": student_id,
                "color": next_auto_color(students),
            }
        else:
            continue

        if student_id in seen:
            continue
        seen.add(student_id)
        price = normalize_amount(raw.get("price"))
        paid_amount = normalize_amount(raw.get("paid_amount")) or 0.0
        members.append({
            "student_id": student_id,
            "name": name or student_id,
            "price": price if price is not None and price >= 0 else "",
            "paid": bool(raw.get("paid", False)),
            "free": bool(raw.get("free", False)),
            "paid_amount": max(0.0, paid_amount),
        })
    return members

def load_settings():
    raw = load_json(SETTINGS_FILE, {})
    raw = raw if isinstance(raw, dict) else {}
    existing_account = bool(raw) or bool(load_json(DATA_FILE)) or bool(load_json(STUDENTS_FILE))
    defaults = {
        "default_reminders_enabled": False,
        "default_student_reminders": False,
        "parent_lesson_end": False,
        "default_send_receipts": True,
        "default_send_receipt_copy": True,
        "zoom_link": "",
        "work_start": "06:00",
        "work_end": "00:00",
        "days_off": [],
        "language": "ru",
        "currency": "RUB",
        "onboarding_completed": False,
        "company_name": "",
        "inn": "",
        "ogrnip": "",
        "address": "",
        "phone": "",
        "service_name": "Услуга",
        "tax_system": "",
        "email_sender": "",
        "thanks_text": "СПАСИБО ЗА ОПЛАТУ!",
        "website": "",
        "bank_name": "",
        "bik": "",
        "account_number": "",
        "corr_account": "",
        "recipient": "",
        "payment_comment": "",
        "student_binding_template": "",
        "parent_binding_template": "",
        "student_reminder_template": "",
        "parent_lesson_end_template": "",
        "teacher_delay_template": "",
        "student_delay_template": "",
        "parent_delay_template": "",
        "receipt_logo": "logo.png" if os.path.exists(os.path.join(current_receipt_assets_dir(), "logo.png")) else "",
        "receipt_signature": "signature.png" if os.path.exists(os.path.join(current_receipt_assets_dir(), "signature.png")) else "",
        "receipt_qrcode": "qrcode.png" if os.path.exists(os.path.join(current_receipt_assets_dir(), "qrcode.png")) else "",
    }
    settings = defaults.copy()
    settings.update(raw)
    # Старые кабинеты без сохранённых рабочих часов сохраняют прежний диапазон.
    if existing_account and "work_start" not in raw:
        settings["work_start"] = "10:00"
    if existing_account and "work_end" not in raw:
        settings["work_end"] = "21:00"
    settings["default_reminders_enabled"] = bool(settings.get("default_reminders_enabled", False))
    settings["default_student_reminders"] = settings.get("default_student_reminders") is True
    settings["parent_lesson_end"] = settings.get("parent_lesson_end") is True
    settings["default_send_receipts"] = bool(settings.get("default_send_receipts", True))
    settings["default_send_receipt_copy"] = bool(settings.get("default_send_receipt_copy", True))
    settings["onboarding_completed"] = bool(settings.get("onboarding_completed", False))
    if settings.get("language") not in {"ru", "en"}:
        settings["language"] = "ru"
    if settings.get("currency") not in {"RUB", "USD", "EUR", "CNY", "TRY"}:
        settings["currency"] = "RUB"
    raw_days_off = settings.get("days_off", [])
    settings["days_off"] = sorted({
        int(day) for day in raw_days_off
        if not isinstance(day, bool) and str(day).isdigit() and 0 <= int(day) <= 6
    }) if isinstance(raw_days_off, list) else []
    return settings


class ReceiptPDF(FPDF):
    def __init__(self):
        super().__init__("P", "mm", (80, 250))
        self.add_font("DejaVu", "", FONT_REGULAR, uni=True)
        self.add_font("DejaVu", "B", FONT_BOLD, uni=True)


def get_receipt_asset_path(settings, key):
    filename = str(settings.get(key, "") or "").strip()
    if not filename:
        return ""
    path = os.path.join(current_receipt_assets_dir(), os.path.basename(filename))
    return path if os.path.exists(path) else ""


def normalize_amount(value):
    try:
        amount = float(str(value).replace(" ", "").replace(",", "."))
        return amount if math.isfinite(amount) else None
    except (TypeError, ValueError):
        return None


def receipt_now():
    try:
        tz = pytz.timezone(TIMEZONE_NAME)
    except pytz.UnknownTimeZoneError:
        tz = pytz.UTC
    return datetime.datetime.now(tz)


def next_receipt_number(now=None):
    now = now or receipt_now()
    return f"{now.strftime('%d%m')}-{now.strftime('%H%M%S')}"


def _receipt_image_size(path):
    """Return image size in pixels for PNG/JPEG using only the standard library."""
    try:
        with open(path, "rb") as f:
            head = f.read(24)
            if head.startswith(b"\x89PNG\r\n\x1a\n") and len(head) >= 24:
                width = int.from_bytes(head[16:20], "big")
                height = int.from_bytes(head[20:24], "big")
                return width, height

            f.seek(0)
            if f.read(2) != b"\xff\xd8":
                return None
            while True:
                marker_start = f.read(1)
                if not marker_start:
                    return None
                if marker_start != b"\xff":
                    continue
                marker = f.read(1)
                while marker == b"\xff":
                    marker = f.read(1)
                if not marker:
                    return None
                marker_code = marker[0]
                if marker_code in (0xD8, 0xD9):
                    continue
                length_bytes = f.read(2)
                if len(length_bytes) != 2:
                    return None
                segment_length = int.from_bytes(length_bytes, "big")
                if segment_length < 2:
                    return None
                if marker_code in {
                    0xC0, 0xC1, 0xC2, 0xC3, 0xC5, 0xC6, 0xC7,
                    0xC9, 0xCA, 0xCB, 0xCD, 0xCE, 0xCF,
                }:
                    precision = f.read(1)
                    height_bytes = f.read(2)
                    width_bytes = f.read(2)
                    if len(precision) != 1 or len(height_bytes) != 2 or len(width_bytes) != 2:
                        return None
                    return int.from_bytes(width_bytes, "big"), int.from_bytes(height_bytes, "big")
                f.seek(segment_length - 2, 1)
    except Exception:
        return None


def _place_receipt_image(pdf, path, box_x, box_y, box_w, box_h):
    """Place an image inside a fixed box without stretching or overlapping neighbours."""
    size = _receipt_image_size(path)
    if size and size[0] > 0 and size[1] > 0:
        source_w, source_h = size
        scale = min(box_w / source_w, box_h / source_h)
        draw_w = source_w * scale
        draw_h = source_h * scale
        draw_x = box_x + (box_w - draw_w) / 2
        draw_y = box_y + (box_h - draw_h) / 2
        pdf.image(path, x=draw_x, y=draw_y, w=draw_w, h=draw_h)
        return draw_h

    # Fallback for an unusual image format: constrain it by height.
    pdf.image(path, x=box_x, y=box_y, h=box_h)
    return box_h


def _receipt_has_any(settings, keys):
    return any(str(settings.get(key, "") or "").strip() for key in keys)


def generate_receipt_pdf(settings, client_name, amount, lesson_id, receipt_number=None, created_at=None, service_name_override=None):
    if REMOTE_STORAGE:
        raise RuntimeError("Бинарное хранилище тестового стенда ещё не подключено.")
    os.makedirs(current_receipts_dir(), exist_ok=True)
    now = created_at or receipt_now()
    receipt_number = receipt_number or next_receipt_number(now)

    pdf = ReceiptPDF()
    pdf.add_page()
    pdf.set_left_margin(5)
    pdf.set_right_margin(5)
    pdf.set_auto_page_break(auto=False)

    logo_path = get_receipt_asset_path(settings, "receipt_logo")
    if logo_path:
        try:
            _place_receipt_image(pdf, logo_path, 25, 5, 30, 20)
            pdf.set_y(27)
        except Exception:
            pdf.set_y(10)
    else:
        pdf.set_y(10)

    pdf.set_font("DejaVu", "B", 12)
    pdf.cell(0, 5, txt="═" * 30, ln=True, align="C")
    pdf.set_font("DejaVu", "B", 14)
    pdf.cell(0, 7, txt="КАССОВЫЙ ЧЕК", ln=True, align="C")
    pdf.set_font("DejaVu", "B", 12)
    pdf.cell(0, 5, txt="═" * 30, ln=True, align="C")
    pdf.ln(2)

    pdf.set_font("DejaVu", "", 9)
    pdf.cell(0, 5, txt=f"Чек №: {receipt_number}", ln=True, align="C")
    pdf.cell(0, 5, txt=now.strftime("%d.%m.%Y %H:%M"), ln=True, align="C")
    pdf.ln(2)

    company_keys = ("company_name", "inn", "ogrnip", "address", "phone")
    if _receipt_has_any(settings, company_keys):
        pdf.set_font("DejaVu", "", 8)
        pdf.cell(0, 4, txt="─" * 30, ln=True, align="C")
        pdf.ln(1)

        company_name = str(settings.get("company_name") or "").strip()
        if company_name:
            pdf.set_font("DejaVu", "B", 10)
            pdf.multi_cell(0, 5, txt=company_name, align="C")
        pdf.set_font("DejaVu", "", 8)
        if settings.get("inn"):
            pdf.cell(0, 4, txt=f"ИНН: {settings['inn']}", ln=True, align="C")
        if settings.get("ogrnip"):
            pdf.cell(0, 4, txt=f"ОГРНИП: {settings['ogrnip']}", ln=True, align="C")
        if settings.get("address"):
            pdf.multi_cell(0, 4, txt=str(settings["address"]), align="C")
        if settings.get("phone"):
            pdf.cell(0, 4, txt=f"Тел.: {settings['phone']}", ln=True, align="C")
        pdf.ln(1)

    pdf.set_font("DejaVu", "", 8)
    pdf.cell(0, 4, txt="─" * 30, ln=True, align="C")
    pdf.ln(2)

    pdf.set_font("DejaVu", "B", 8)
    pdf.cell(40, 5, "Наименование", border=1, align="C")
    pdf.cell(20, 5, "Цена", border=1, align="C")
    pdf.cell(10, 5, "Кол", border=1, align="C")
    pdf.ln()

    service_name = str(service_name_override or settings.get("service_name") or "Услуга")
    pdf.set_font("DejaVu", "", 8)
    pdf.cell(40, 5, service_name[:28], border=1, align="L")
    pdf.cell(20, 5, f"{amount:.2f}", border=1, align="R")
    pdf.cell(10, 5, "1", border=1, align="C")
    pdf.ln()
    pdf.ln(2)

    pdf.set_font("DejaVu", "", 9)
    pdf.cell(50, 5, txt="СУММА БЕЗ НДС", border=0)
    pdf.cell(20, 5, txt=f"{amount:.2f}", border=0, align="R")
    pdf.ln()
    pdf.set_font("DejaVu", "B", 10)
    pdf.cell(50, 6, txt="ИТОГО:", border=0)
    pdf.cell(20, 6, txt=f"{amount:.2f}", border=0, align="R")
    pdf.ln()
    pdf.set_font("DejaVu", "", 8)
    pdf.cell(50, 4, txt="Безналичными", border=0)
    pdf.cell(20, 4, txt=f"{amount:.2f}", border=0, align="R")
    pdf.ln(3)

    tax_system = str(settings.get("tax_system") or "").strip()
    if tax_system:
        pdf.cell(0, 4, txt="─" * 30, ln=True, align="C")
        pdf.ln(1)
        pdf.cell(0, 4, txt=f"Система: {tax_system}", ln=True, align="C")
        pdf.ln(1)

    requisites_keys = ("bank_name", "bik", "account_number", "corr_account", "recipient")
    if _receipt_has_any(settings, requisites_keys):
        pdf.cell(0, 4, txt="─" * 30, ln=True, align="C")
        pdf.ln(1)
        pdf.set_font("DejaVu", "B", 8)
        pdf.cell(0, 4, txt="РЕКВИЗИТЫ:", ln=True, align="C")
        pdf.set_font("DejaVu", "", 7)
        if settings.get("bank_name"):
            pdf.multi_cell(0, 3, txt=str(settings["bank_name"]), align="C")
        if settings.get("bik"):
            pdf.cell(0, 3, txt=f"БИК: {settings['bik']}", ln=True, align="C")
        if settings.get("account_number"):
            pdf.cell(0, 3, txt=f"Сч: {settings['account_number']}", ln=True, align="C")
        if settings.get("corr_account"):
            pdf.cell(0, 3, txt=f"Корр. сч: {settings['corr_account']}", ln=True, align="C")
        if settings.get("recipient"):
            pdf.multi_cell(0, 3, txt=f"Получатель: {settings['recipient']}", align="C")
        pdf.ln(1)

    contact_keys = ("email_sender", "website")
    if _receipt_has_any(settings, contact_keys):
        pdf.cell(0, 4, txt="─" * 30, ln=True, align="C")
        pdf.ln(1)
        pdf.set_font("DejaVu", "", 7)
        if settings.get("email_sender"):
            pdf.cell(0, 4, txt=f"Email: {settings['email_sender']}", ln=True, align="C")
        if settings.get("website"):
            pdf.cell(0, 4, txt=str(settings["website"]), ln=True, align="C")

    thanks_text = str(settings.get("thanks_text") or "").strip()
    if thanks_text:
        pdf.ln(2)
        pdf.set_font("DejaVu", "B", 8)
        pdf.multi_cell(0, 4, txt=thanks_text, align="C")

    # Footer assets live in separate reserved boxes. Their position is based on
    # the content end, but never allowed to overlap each other or the page edge.
    signature_path = get_receipt_asset_path(settings, "receipt_signature")
    qr_path = get_receipt_asset_path(settings, "receipt_qrcode")

    footer_y = max(pdf.get_y() + 4, 165)
    signature_box_h = 20 if signature_path else 0
    qr_box_h = 24 if qr_path else 0
    gap = 5 if signature_path and qr_path else 0
    footer_height = signature_box_h + gap + qr_box_h

    if footer_y + footer_height > 242:
        pdf.add_page()
        footer_y = 12

    if signature_path:
        try:
            _place_receipt_image(pdf, signature_path, 10, footer_y, 60, signature_box_h)
        except Exception:
            pass
        footer_y += signature_box_h

    if signature_path and qr_path:
        footer_y += gap

    if qr_path:
        try:
            _place_receipt_image(pdf, qr_path, 28, footer_y, 24, qr_box_h)
        except Exception:
            pass
        footer_y += qr_box_h

    pdf.set_y(min(footer_y + 2, 245))

    safe_lesson_id = re.sub(r"[^A-Za-z0-9_-]", "_", str(lesson_id))
    filename = f"check_{safe_lesson_id}_{receipt_number.replace('-', '_')}.pdf"
    path = os.path.join(current_receipts_dir(), filename)
    pdf.output(path)
    return path, receipt_number, now


@serialized_data
def init_book():
    if REMOTE_STORAGE:
        return
    book_file = current_book_file()
    if os.path.exists(book_file):
        return
    wb = openpyxl.Workbook()
    try:
        ws = wb.active
        ws.title = "Книга учёта"
        headers = ["№ п/п", "Дата и время", "№ Квитанции", "Клиент", "доходы, руб", "Статус"]
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = Font(bold=True)
            cell.alignment = Alignment(horizontal="center")
            cell.border = Border(left=Side(style="thin"), right=Side(style="thin"), top=Side(style="thin"), bottom=Side(style="thin"))
        ws.column_dimensions["A"].width = 8
        ws.column_dimensions["B"].width = 18
        ws.column_dimensions["C"].width = 20
        ws.column_dimensions["D"].width = 20
        ws.column_dimensions["E"].width = 15
        ws.column_dimensions["F"].width = 12
        fd, temporary = tempfile.mkstemp(prefix="book_", suffix=".xlsx", dir=os.path.dirname(book_file))
        os.close(fd)
        try:
            wb.save(temporary)
            os.replace(temporary, book_file)
        finally:
            if os.path.exists(temporary):
                os.remove(temporary)
    finally:
        wb.close()


def add_receipt_to_book(client_name, amount, receipt_number, created_at, status="Оплачено"):
    with DATA_LOCK:
        init_book()
        book_file = current_book_file()
        # Загружаем книгу из памяти, чтобы openpyxl/ZipFile не удерживал блокировку
        # исходного book.xlsx во время атомарной замены или отката транзакции.
        with open(book_file, "rb") as source:
            book_data = source.read()
        wb = openpyxl.load_workbook(io.BytesIO(book_data))
        fd, temp_book = tempfile.mkstemp(prefix="book_", suffix=".xlsx", dir=os.path.dirname(book_file) or BASE_DIR)
        os.close(fd)
        try:
            try:
                ws = wb.active
                row_num = ws.max_row + 1
                ws.cell(row=row_num, column=1, value=row_num - 1)
                ws.cell(row=row_num, column=2, value=created_at.strftime("%d.%m.%Y %H:%M"))
                ws.cell(row=row_num, column=3, value=receipt_number)
                ws.cell(row=row_num, column=4, value=client_name)
                ws.cell(row=row_num, column=5, value=amount)
                ws.cell(row=row_num, column=6, value=status)
                for col in range(1, 7):
                    ws.cell(row=row_num, column=col).alignment = Alignment(horizontal="center")
                wb.save(temp_book)
            finally:
                wb.close()
            os.replace(temp_book, book_file)
        finally:
            if os.path.exists(temp_book):
                os.remove(temp_book)


def _atomic_copy(source_path, destination_path):
    directory = os.path.dirname(os.path.abspath(destination_path)) or BASE_DIR
    os.makedirs(directory, exist_ok=True)
    fd, temp_path = tempfile.mkstemp(prefix="payment_copy_", suffix=".tmp", dir=directory)
    try:
        with os.fdopen(fd, "wb") as target, open(source_path, "rb") as source:
            shutil.copyfileobj(source, target)
            target.flush()
            os.fsync(target.fileno())
        os.replace(temp_path, destination_path)
    finally:
        if os.path.exists(temp_path):
            os.remove(temp_path)


def payment_transaction_paths():
    root = tenant_root()
    return {
        "marker": os.path.join(root, ".payment_transaction.json"),
        "schedule_backup": os.path.join(root, ".payment_schedule.before"),
        "book_backup": os.path.join(root, ".payment_book.before"),
        "payments_backup": os.path.join(root, ".payment_payments.before"),
    }


def payments_file():
    return os.path.join(tenant_root(), "payments.json")


def _cleanup_payment_transaction(paths):
    for key in ("marker", "schedule_backup", "book_backup", "payments_backup"):
        path = paths[key]
        if os.path.exists(path):
            os.remove(path)


def recover_payment_transaction():
    """Recover a payment interrupted between schedule.json and book.xlsx writes."""
    paths = payment_transaction_paths()
    marker_path = paths["marker"]
    with DATA_LOCK:
        # Another request may have finished the transaction while we waited.
        if not os.path.exists(marker_path):
            return False
        try:
            with open(marker_path, "r", encoding="utf-8") as source:
                marker = json.load(source)
        except (OSError, json.JSONDecodeError) as exc:
            raise DataCorruptionError("Повреждён журнал платёжной транзакции; автоматическое восстановление остановлено.") from exc

        schedule_file = tenant_file(DATA_FILE)
        book_file = current_book_file()
        for existed_key, backup_key, target in (
            ("schedule_existed", "schedule_backup", schedule_file),
            ("book_existed", "book_backup", book_file),
            ("payments_existed", "payments_backup", payments_file()),
        ):
            if existed_key not in marker:  # Journals created before payment history existed.
                continue
            if marker.get(existed_key):
                backup = paths[backup_key]
                if not os.path.exists(backup):
                    raise DataCorruptionError("Не найдена резервная копия незавершённой платёжной транзакции.")
                _atomic_copy(backup, target)
            elif os.path.exists(target):
                os.remove(target)
        _cleanup_payment_transaction(paths)
    return True


@contextmanager
def payment_files_transaction():
    """Atomically coordinate schedule and book writes with crash recovery."""
    if REMOTE_STORAGE:
        raise RuntimeError("Оплаты отключены до подключения удалённого бинарного хранилища.")
    with DATA_LOCK:
        recover_payment_transaction()
        paths = payment_transaction_paths()
        schedule_file = tenant_file(DATA_FILE)
        book_file = current_book_file()
        marker = {
            "schedule_existed": os.path.exists(schedule_file),
            "book_existed": os.path.exists(book_file),
            "payments_existed": os.path.exists(payments_file()),
            "created_at": datetime.datetime.utcnow().replace(microsecond=0).isoformat() + "Z",
        }
        if marker["schedule_existed"]:
            _atomic_copy(schedule_file, paths["schedule_backup"])
        if marker["book_existed"]:
            _atomic_copy(book_file, paths["book_backup"])
        if marker["payments_existed"]:
            _atomic_copy(payments_file(), paths["payments_backup"])
        _save_json_raw(paths["marker"], marker)
        try:
            yield
        except BaseException:
            recover_payment_transaction()
            raise
        else:
            _cleanup_payment_transaction(paths)


def numeric_telegram_chat_id(value):
    value = str(value or "").strip()
    if re.fullmatch(r"-?\d+", value):
        return int(value)
    return None


async def send_receipt_document(chat_id, pdf_path, caption, filename=None):
    if BOT_APPLICATION is None:
        raise RuntimeError("Telegram bot application ещё не готов")
    with open(pdf_path, "rb") as receipt_file:
        await BOT_APPLICATION.bot.send_document(
            chat_id=chat_id,
            document=receipt_file,
            filename=filename or os.path.basename(pdf_path),
            caption=caption,
        )


def send_receipt_from_flask(chat_id, pdf_path, caption, filename=None):
    if BOT_LOOP is None or BOT_APPLICATION is None:
        return False, "Telegram-бот ещё не готов к отправке."
    future = asyncio.run_coroutine_threadsafe(send_receipt_document(chat_id, pdf_path, caption, filename), BOT_LOOP)
    try:
        future.result(timeout=25)
        return True, ""
    except Exception as exc:
        return False, str(exc)


def delete_receipt_file(pdf_path):
    """Удаляет временный PDF чека после отправки/обработки."""
    if not pdf_path:
        return
    try:
        path = os.path.abspath(str(pdf_path))
        receipts_root = os.path.abspath(current_receipts_dir())
        if os.path.commonpath([path, receipts_root]) != receipts_root:
            return
        if os.path.isfile(path):
            os.remove(path)
    except Exception as exc:
        print(f"Не удалось удалить временный чек {pdf_path}: {exc}")


def validate_init_data(init_data):
    if ALLOW_UNAUTHENTICATED:
        return True, None
    if not TOKEN or not init_data:
        return False, None

    try:
        parsed = dict(parse_qsl(init_data, keep_blank_values=True))
        received_hash = parsed.pop("hash", "")
        if not received_hash:
            return False, None

        auth_date = int(parsed.get("auth_date", "0"))
        if auth_date <= 0 or abs(int(time.time()) - auth_date) > 86400:
            return False, None

        data_check_string = "\n".join(f"{key}={parsed[key]}" for key in sorted(parsed))
        secret_key = hmac.new(b"WebAppData", TOKEN.encode("utf-8"), hashlib.sha256).digest()
        calculated_hash = hmac.new(secret_key, data_check_string.encode("utf-8"), hashlib.sha256).hexdigest()
        if not hmac.compare_digest(calculated_hash, received_hash):
            return False, None

        user = json.loads(parsed.get("user", "{}")) if parsed.get("user") else {}
        # Teacher identity must always come from the signed Telegram user.  A
        # signed payload without a user must never inherit OWNER_ID.
        if not isinstance(user, dict):
            return False, None
        user_id = user.get("id")
        if isinstance(user_id, bool) or not isinstance(user_id, int) or user_id <= 0:
            return False, None
        return True, user
    except (ValueError, TypeError, json.JSONDecodeError):
        return False, None


def current_series_key(lesson):
    return lesson.get("series_id") or lesson.get("id")


def same_series(lesson, series_key):
    return (lesson.get("series_id") or lesson.get("id")) == series_key


def is_personal_event(item):
    return isinstance(item, dict) and item.get("entry_type") == "personal"


def next_lesson_id():
    return f"l_{time.time_ns()}"


def next_series_id():
    return f"s_{time.time_ns()}"


flask_app = Flask(__name__)
CORS(
    flask_app,
    resources={r"/api/*": {"origins": [WEBAPP_ORIGIN]}},
    allow_headers=["Content-Type", "X-Telegram-Init-Data"],
    methods=["GET", "POST", "OPTIONS"],
)


@flask_app.errorhandler(DataCorruptionError)
def handle_data_corruption(error):
    return jsonify({"status": "error", "message": str(error)}), 503


@flask_app.before_request
def protect_api():
    if request.method == "OPTIONS" or request.path == "/api/health":
        return None
    if not request.path.startswith("/api/"):
        return None

    ok, user = validate_init_data(request.headers.get("X-Telegram-Init-Data", ""))
    if not ok:
        return jsonify({"status": "error", "message": "Недействительные данные Telegram WebApp."}), 401
    g.telegram_user = user or {}
    if ALLOW_UNAUTHENTICATED and not user:
        teacher_id = str(OWNER_ID or "local-dev").strip()
    else:
        teacher_id = str((user or {}).get("id", "")).strip()
    if not teacher_id:
        return jsonify({"status": "error", "message": "Не удалось определить преподавателя Telegram."}), 401
    from invitation_channels import recipient_only
    if recipient_only(sys.modules[__name__], teacher_id):
        return jsonify(status='error', code='recipient_only'), 403
    ensure_teacher_registered(teacher_id, user or {})
    TEACHER_CONTEXT.set(teacher_id)
    g.teacher_id = teacher_id
    recover_payment_transaction()
    return None






def get_student_record(students, student_id):
    raw = students.get(str(student_id), {})
    if isinstance(raw, dict):
        return raw
    if isinstance(raw, str):
        return {"name": raw}
    return {}


def get_student_price(students, student_id, fallback=None):
    # Новая модель: цена хранится у конкретного занятия/участника группы.
    # default_price оставлен только как совместимость со старыми данными.
    fallback_price = normalize_amount(fallback)
    if fallback_price is not None and fallback_price > 0:
        return fallback_price
    info = get_student_record(students, student_id)
    legacy_price = normalize_amount(info.get("default_price"))
    return legacy_price if legacy_price is not None and legacy_price > 0 else None


def member_lesson_price(students, member, lesson):
    return get_student_price(students, member.get("student_id"), member.get("price", lesson.get("price")))


def lesson_price_for_student(students, lesson, student_id):
    student_id = str(student_id)
    if lesson.get("lesson_type") == "group":
        for member in lesson.get("group_members") or []:
            if str(member.get("student_id", "")) == student_id:
                return member_lesson_price(students, member, lesson)
        return None
    if str(lesson.get("student_id", "")) != student_id:
        return None
    return get_student_price(students, student_id, lesson.get("price"))


def next_auto_color(students):
    used = []
    for raw in students.values():
        if isinstance(raw, dict):
            try:
                c = int(raw.get("color"))
            except (TypeError, ValueError):
                continue
            if 0 <= c <= 6:
                used.append(c)
    for c in range(7):
        if c not in used:
            return c
    return len(used) % 7


def student_lesson_stats(student_id, schedule=None):
    student_id = str(student_id)
    schedule = schedule if schedule is not None else load_json(DATA_FILE)
    now = receipt_now()
    paid_lessons = 0
    conducted_lessons = 0
    for date_key, lessons in schedule.items():
        for lesson in lessons:
            if lesson.get("cancelled"):
                continue
            involved = False
            paid = False
            if lesson.get("lesson_type") == "group":
                for member in lesson.get("group_members") or []:
                    if str(member.get("student_id", "")) == student_id:
                        involved = True
                        paid = bool(member.get("paid")) or bool(member.get("free"))
                        break
            elif str(lesson.get("student_id", "")) == student_id:
                involved = True
                paid = bool(lesson.get("paid")) or bool(lesson.get("free"))
            if not involved:
                continue
            if paid:
                paid_lessons += 1
            try:
                lesson_dt = datetime.datetime.strptime(f"{date_key} {lesson.get('time', '00:00')}", "%Y-%m-%d %H:%M")
                if now.tzinfo is not None:
                    lesson_dt = now.tzinfo.localize(lesson_dt) if hasattr(now.tzinfo, 'localize') else lesson_dt.replace(tzinfo=now.tzinfo)
                if lesson_dt <= now:
                    conducted_lessons += 1
            except (TypeError, ValueError):
                pass
    history = []
    total_lessons = 0
    for date_key, lessons in schedule.items():
        for lesson in lessons:
            if lesson.get("cancelled"):
                continue
            involved = False
            paid = False
            if lesson.get("lesson_type") == "group":
                for member in lesson.get("group_members") or []:
                    if str(member.get("student_id", "")) == student_id:
                        involved = True
                        paid = bool(member.get("paid")) or bool(member.get("free"))
                        break
            elif str(lesson.get("student_id", "")) == student_id:
                involved = True
                paid = bool(lesson.get("paid")) or bool(lesson.get("free"))
            if not involved:
                continue
            total_lessons += 1
            try:
                item_dt = datetime.datetime.strptime(f"{date_key} {lesson.get('time', '00:00')}", "%Y-%m-%d %H:%M")
                if now.tzinfo is not None:
                    item_dt = now.tzinfo.localize(item_dt) if hasattr(now.tzinfo, 'localize') else item_dt.replace(tzinfo=now.tzinfo)
                is_conducted = item_dt <= now
            except (TypeError, ValueError):
                is_conducted = False
            if is_conducted:
                history.append({
                    "date": date_key,
                    "time": str(lesson.get("time", "")),
                    "paid": paid,
                    "report": str(lesson.get("report", "") or ""),
                })
    history.sort(key=lambda item: (item.get("date", ""), item.get("time", "")), reverse=True)
    return {
        "paid_lessons": paid_lessons,
        "conducted_lessons": conducted_lessons,
        "total_lessons": total_lessons,
        "history": history[:8],
    }


def safe_float(value, default=0.0):
    try:
        return float(str(value).replace(",", "."))
    except (TypeError, ValueError):
        return default


def parse_minutes(value):
    try:
        h, m = str(value).split(":", 1)
        return int(h) * 60 + int(m)
    except Exception:
        return 0


def format_minutes(value):
    value = max(0, min(24 * 60, int(value)))
    if value == 24 * 60:
        return "00:00"
    return f"{value // 60:02d}:{value % 60:02d}"


def parse_work_end_minutes(value):
    return 24 * 60 if str(value).strip() == "00:00" else parse_minutes(value)


def build_work_center(week_start):
    try:
        monday = datetime.datetime.strptime(week_start, "%Y-%m-%d").date()
    except (TypeError, ValueError):
        today_local = receipt_now().date()
        monday = today_local - datetime.timedelta(days=today_local.weekday())

    schedule = load_json(DATA_FILE)
    students = load_json(STUDENTS_FILE)
    settings = load_settings()
    now = receipt_now()
    today = now.date()

    attention = []
    teacher_id = str(getattr(g, "teacher_id", "") or "").strip()
    if teacher_id:
        students.pop(teacher_id, None)
    paid_sum = 0.0
    planned_sum = 0.0
    lesson_count = 0
    paid_count = 0
    windows = []
    days_off = set(settings.get("days_off", []))

    for i in range(7):
        day = monday + datetime.timedelta(days=i)
        key = day.strftime("%Y-%m-%d")
        day_lessons = sorted([x for x in schedule.get(key, []) if not x.get("cancelled")], key=lambda x: x.get("time", "00:00"))
        lesson_count += sum(1 for lesson in day_lessons if not is_personal_event(lesson))
        occupied = []
        for lesson in day_lessons:
            if is_personal_event(lesson):
                pass
            elif lesson.get("lesson_type") == "group":
                member_amounts = [(member, member_lesson_price(students, member, lesson) or 0.0) for member in (lesson.get("group_members") or []) if not member.get("free")]
                planned_sum += sum(amount for _member, amount in member_amounts)
                paid_sum += sum(amount for member, amount in member_amounts if member.get("paid"))
                if lesson.get("paid"):
                    paid_count += 1
            else:
                amount = get_student_price(students, lesson.get("student_id"), lesson.get("price")) or 0.0
                planned_sum += amount
                if lesson.get("paid"):
                    paid_count += 1
                    paid_sum += amount
            start = parse_minutes(lesson.get("time", "06:00"))
            duration = max(5, int(lesson.get("duration", 60) or 60))
            occupied.append((start, min(24 * 60, start + duration)))
        if i in days_off:
            continue
        occupied.sort()
        merged = []
        for start, end in occupied:
            if not merged or start > merged[-1][1]:
                merged.append([start, end])
            else:
                merged[-1][1] = max(merged[-1][1], end)
        work_start = parse_minutes(settings.get("work_start", "06:00"))
        work_end = parse_work_end_minutes(settings.get("work_end", "00:00"))
        if work_end <= work_start:
            work_start, work_end = 6 * 60, 24 * 60
        cursor = work_start
        for start, end in merged:
            clipped_start = max(work_start, start)
            clipped_end = min(work_end, end)
            if clipped_end <= work_start or clipped_start >= work_end:
                continue
            if clipped_start - cursor >= 60:
                windows.append({"date": key, "from": format_minutes(cursor), "to": format_minutes(clipped_start)})
            cursor = max(cursor, clipped_end)
        if work_end - cursor >= 60:
            windows.append({"date": key, "from": format_minutes(cursor), "to": format_minutes(work_end)})

    debt_map = {}
    next_candidates = []
    for date_key, lessons in schedule.items():
        try:
            lesson_date = datetime.datetime.strptime(date_key, "%Y-%m-%d").date()
        except ValueError:
            continue
        for lesson in lessons:
            if lesson.get("cancelled"):
                continue
            try:
                start_naive = datetime.datetime.strptime(f"{date_key} {lesson.get('time', '00:00')}", "%Y-%m-%d %H:%M")
                if hasattr(now.tzinfo, "localize"):
                    lesson_start = now.tzinfo.localize(start_naive)
                else:
                    lesson_start = start_naive.replace(tzinfo=now.tzinfo)
            except (TypeError, ValueError):
                continue
            duration = max(5, int(lesson.get("duration", 60) or 60))
            lesson_end = lesson_start + datetime.timedelta(minutes=duration)

            if is_personal_event(lesson):
                continue

            if lesson_end >= now:
                is_now = lesson_start <= now < lesson_end
                next_candidates.append((0 if is_now else 1, lesson_start, lesson_end, date_key, lesson))

            if lesson_start <= now:
                if lesson.get("lesson_type") == "group":
                    for member in lesson.get("group_members") or []:
                        if member.get("paid") or member.get("free"):
                            continue
                        sid = str(member.get("student_id", "")).strip()
                        name = str(member.get("name", "") or get_student_record(students, sid).get("name", "Ученик"))
                        amount = max(0.0, (member_lesson_price(students, member, lesson) or 0.0) - (normalize_amount(member.get("paid_amount")) or 0.0))
                        item = debt_map.setdefault(sid or name, {"student_id": sid, "name": name, "unpaid_count": 0, "amount": 0.0})
                        item["unpaid_count"] += 1
                        item["amount"] += amount
                elif not lesson.get("paid") and not lesson.get("free"):
                    sid = str(lesson.get("student_id", "")).strip()
                    name = str(lesson.get("student", "") or get_student_record(students, sid).get("name", "Ученик"))
                    amount = max(0.0, (get_student_price(students, sid, lesson.get("price")) or 0.0) - (normalize_amount(lesson.get("paid_amount")) or 0.0))
                    item = debt_map.setdefault(sid or name, {"student_id": sid, "name": name, "unpaid_count": 0, "amount": 0.0})
                    item["unpaid_count"] += 1
                    item["amount"] += amount

            if lesson_date < today and not lesson.get("paid") and (get_student_price(students, lesson.get("student_id"), lesson.get("price")) if lesson.get("lesson_type") != "group" else True):
                attention.append({
                    "type": "unpaid",
                    "date": date_key,
                    "student": lesson.get("student", "Ученик"),
                    "time": lesson.get("time", ""),
                    "text": f"Не оплачено: {lesson.get('student', 'Ученик')} · {date_key} {lesson.get('time', '')}",
                })

    current_lesson = None
    next_lesson = None

    def _lesson_payload(lesson_start, lesson_end, date_key, lesson, status):
        is_group = lesson.get("lesson_type") == "group"
        student_id = "" if is_group else str(lesson.get("student_id", "")).strip()
        student_info = get_student_record(students, student_id) if student_id else {}
        individual_zoom = str(student_info.get("zoom_link", "") or "").strip()
        legacy_lesson_zoom = str(lesson.get("zoom_link", "") or "").strip()
        common_zoom = str(settings.get("zoom_link", "") or "").strip()
        return {
            "date": date_key,
            "id": lesson.get('id'),
            "time": str(lesson.get("time", "")),
            "end_time": lesson_end.strftime("%H:%M"),
            "starts_at": lesson_start.isoformat(),
            "ends_at": lesson_end.isoformat(),
            "group_members": [{"student_id": str(m.get("student_id", "")), "name": str(m.get("name", ""))}
                              for m in lesson.get("group_members") or []] if is_group else [],
            "duration": int(lesson.get("duration", 60) or 60),
            "student": str(lesson.get("group_name") or lesson.get("student") or "Ученик"),
            "student_id": student_id,
            "lesson_type": "group" if is_group else "student",
            "status": status,
            "minutes_until": max(0, int((lesson_start - now).total_seconds() // 60)),
            "board_link": "" if is_group else str(student_info.get("board_link", "") or "").strip(),
            "zoom_link": individual_zoom or legacy_lesson_zoom or common_zoom,
        }

    current_rows = []
    future_rows = []
    for _rank, lesson_start, lesson_end, date_key, lesson in next_candidates:
        if lesson_start <= now < lesson_end:
            current_rows.append((lesson_start, lesson_end, date_key, lesson))
        elif lesson_start > now:
            future_rows.append((lesson_start, lesson_end, date_key, lesson))

    if current_rows:
        current_rows.sort(key=lambda row: row[0])
        lesson_start, lesson_end, date_key, lesson = current_rows[0]
        current_lesson = _lesson_payload(lesson_start, lesson_end, date_key, lesson, "now")
    if future_rows:
        future_rows.sort(key=lambda row: row[0])
        lesson_start, lesson_end, date_key, lesson = future_rows[0]
        next_lesson = _lesson_payload(lesson_start, lesson_end, date_key, lesson, "next")
    elif not current_lesson and next_candidates:
        next_candidates.sort(key=lambda item: item[1])
        _rank, lesson_start, lesson_end, date_key, lesson = next_candidates[0]
        next_lesson = _lesson_payload(lesson_start, lesson_end, date_key, lesson, "next")

    birthdays = []
    for student_id, raw in students.items():
        if not isinstance(raw, dict):
            continue
        birthday = str(raw.get("birthday", "")).strip()
        if birthday:
            try:
                source = datetime.datetime.strptime(birthday, "%Y-%m-%d").date()
                candidate = datetime.date(today.year, source.month, source.day)
                if candidate < today:
                    candidate = datetime.date(today.year + 1, source.month, source.day)
                days = (candidate - today).days
                if days <= 30:
                    item = {"student_id": student_id, "name": raw.get("name", "Ученик"), "date": candidate.strftime("%Y-%m-%d"), "days": days}
                    birthdays.append(item)
                    if days <= 7:
                        attention.append({"type": "birthday", "student_id": student_id, "text": f"День рождения: {item['name']} · через {days} дн." if days else f"Сегодня день рождения: {item['name']}"})
            except ValueError:
                pass
        balance = safe_float(raw.get("balance", 0), 0.0)
        if balance < 0:
            attention.append({"type": "balance", "student_id": student_id, "text": f"Долг: {raw.get('name', 'Ученик')} · {abs(balance):.0f} ₽"})

    # Блок "Требует внимания" удалён из интерфейса: не формируем дублирующие записи.
    attention = []
    birthdays.sort(key=lambda x: x["days"])
    debts = list(debt_map.values())
    for item in debts:
        item["amount"] = round(item["amount"], 2)
    debts.sort(key=lambda item: (-item["amount"], str(item["name"]).lower()))
    debt_total = round(sum(item["amount"] for item in debts), 2)

    return {
        "attention": attention,
        "debts": debts,
        "debt_total": debt_total,
        "current_lesson": current_lesson,
        "next_lesson": next_lesson,
        "windows": windows,
        "birthdays": birthdays,
        "summary": {
            "week_start": monday.strftime("%Y-%m-%d"),
            "lessons": lesson_count,
            "paid": paid_count,
            "unpaid": max(0, lesson_count - paid_count),
            "planned_sum": round(planned_sum, 2),
            "paid_sum": round(paid_sum, 2),
        },
    }


RUS_WEEKDAYS = ["Понедельник", "Вторник", "Среда", "Четверг", "Пятница", "Суббота", "Воскресенье"]
RUS_MONTHS = ["января", "февраля", "марта", "апреля", "мая", "июня", "июля", "августа", "сентября", "октября", "ноября", "декабря"]


def format_ru_date(value):
    return f"{value.day} {RUS_MONTHS[value.month - 1]} {value.year}"


def lesson_end_time(lesson):
    try:
        start = datetime.datetime.strptime(str(lesson.get("time", "00:00")), "%H:%M")
        duration = max(1, int(lesson.get("duration", 60) or 60))
        return (start + datetime.timedelta(minutes=duration)).strftime("%H:%M")
    except (TypeError, ValueError):
        return ""


def pdf_wrap_lines(pdf, text, width):
    text = str(text or "").strip()
    if not text:
        return []
    words = text.split()
    lines = []
    current = ""
    for word in words:
        candidate = word if not current else f"{current} {word}"
        if pdf.get_string_width(candidate) <= width:
            current = candidate
            continue
        if current:
            lines.append(current)
        # Very long tokens (for example links) are split safely.
        current = ""
        chunk = ""
        for ch in word:
            test = chunk + ch
            if chunk and pdf.get_string_width(test) > width:
                lines.append(chunk)
                chunk = ch
            else:
                chunk = test
        current = chunk
    if current:
        lines.append(current)
    return lines


def weekly_lesson_lines(pdf, lesson, text_width, students):
    start = str(lesson.get("time", ""))
    end = lesson_end_time(lesson)
    duration_text = f"{start}-{end}" if end else start
    if lesson.get("lesson_type") == "group":
        name = str(lesson.get("group_name") or lesson.get("student") or "Группа")
    else:
        student_id = str(lesson.get("student_id", ""))
        student_info = students.get(student_id, {}) if isinstance(students, dict) else {}
        name = str((student_info or {}).get("calendar_name") or lesson.get("student") or (student_info or {}).get("name") or "Ученик")
    pdf.set_font("DejaVu", "B", 7)
    return pdf_wrap_lines(pdf, f"{duration_text}  {name}", text_width)


def build_week_schedule_pdf(start_date, week_data):
    students = load_json(STUDENTS_FILE)
    if not os.path.exists(FONT_REGULAR) or not os.path.exists(FONT_BOLD):
        raise RuntimeError("Не найдены шрифты DejaVu для формирования PDF.")

    pdf = FPDF("L", "mm", "A4")
    pdf.set_auto_page_break(False)
    pdf.add_font("DejaVu", "", FONT_REGULAR, uni=True)
    pdf.add_font("DejaVu", "B", FONT_BOLD, uni=True)

    page_w, page_h = 297, 210
    margin_x = 8
    bottom_margin = 8
    usable_w = page_w - margin_x * 2
    col_w = usable_w / 7
    text_pad = 2
    header_y = 22
    header_h = 13
    content_y = header_y + header_h + 2
    content_bottom = page_h - bottom_margin
    line_h = 3.8

    days = []
    for i in range(7):
        day = start_date + datetime.timedelta(days=i)
        key = day.strftime("%Y-%m-%d")
        days.append((day, list(week_data.get(key, []))))

    positions = [0] * 7
    page_no = 0
    while True:
        page_no += 1
        pdf.add_page()
        week_end = start_date + datetime.timedelta(days=6)
        pdf.set_font("DejaVu", "B", 14)
        pdf.cell(0, 7, f"TEMLI — расписание недели: {format_ru_date(start_date)} - {format_ru_date(week_end)}", ln=True, align="C")
        if page_no > 1:
            pdf.set_font("DejaVu", "", 8)
            pdf.cell(0, 4, f"Продолжение · страница {page_no}", ln=True, align="C")

        # Day headers.
        for i, (day, _) in enumerate(days):
            x = margin_x + i * col_w
            pdf.set_xy(x, header_y)
            pdf.set_font("DejaVu", "B", 8)
            pdf.cell(col_w, 6, RUS_WEEKDAYS[i], border=1, ln=0, align="C")
            pdf.set_xy(x, header_y + 6)
            pdf.set_font("DejaVu", "", 7.5)
            pdf.cell(col_w, 7, day.strftime("%d.%m"), border=1, ln=0, align="C")

        any_remaining = False
        made_progress = False
        for i, (_, lessons) in enumerate(days):
            x = margin_x + i * col_w
            y = content_y
            idx = positions[i]
            if idx >= len(lessons):
                if page_no == 1 and not lessons:
                    pdf.set_xy(x + text_pad, y + 2)
                    pdf.set_font("DejaVu", "", 7)
                    pdf.cell(col_w - text_pad * 2, 4, "Нет занятий", align="C")
                continue

            while idx < len(lessons):
                lesson = lessons[idx]
                pdf.set_font("DejaVu", "", 7)
                lines = weekly_lesson_lines(pdf, lesson, col_w - text_pad * 2, students)
                block_h = max(11, len(lines) * line_h + 4)
                if y + block_h > content_bottom:
                    any_remaining = True
                    break

                pdf.rect(x, y, col_w, block_h)
                pdf.set_xy(x + text_pad, y + 2)
                for line_no, line in enumerate(lines):
                    pdf.set_x(x + text_pad)
                    if line_no == 0:
                        pdf.set_font("DejaVu", "B", 7)
                    else:
                        pdf.set_font("DejaVu", "", 6.7)
                    pdf.cell(col_w - text_pad * 2, line_h, line, ln=1)
                y += block_h + 1.5
                idx += 1
                positions[i] = idx
                made_progress = True

            if idx < len(lessons):
                any_remaining = True

        if not any_remaining:
            break
        if not made_progress:
            # Safety guard: a pathological single lesson should never loop forever.
            break

    raw = pdf.output(dest="S")
    if isinstance(raw, str):
        raw = raw.encode("latin-1")
    return io.BytesIO(raw)


@flask_app.route("/api/download_book", methods=["GET"])
def download_book():
    if REMOTE_STORAGE:
        return jsonify({
            "status": "error",
            "message": "Книга учёта временно отключена на тестовом стенде удалённого хранения.",
        }), 503
    request_user = getattr(g, "telegram_user", {}) or {}
    teacher_target = str(getattr(g, "teacher_id", "") or request_user.get("id", "")).strip()
    teacher_chat_id = numeric_telegram_chat_id(teacher_target)
    if teacher_chat_id is None:
        return jsonify({"status": "error", "message": "Не удалось определить ваш Telegram ID. Откройте приложение через бота."}), 400

    temp_path = None
    try:
        with DATA_LOCK:
            init_book()
            with open(current_book_file(), "rb") as source:
                data = source.read()
        fd, temp_path = tempfile.mkstemp(prefix="book_export_", suffix=".xlsx")
        with os.fdopen(fd, "wb") as temp_file:
            temp_file.write(data)
        ok, error = send_receipt_from_flask(
            teacher_chat_id,
            temp_path,
            "📗 TEMLI · Книга учёта",
            filename="temli_kniga_ucheta.xlsx",
        )
        if not ok:
            return jsonify({"status": "error", "message": f"Не удалось отправить книгу в Telegram: {error}"}), 502
        return jsonify({"status": "ok", "message": "Книга учёта отправлена вам в Telegram."})
    finally:
        if temp_path and os.path.exists(temp_path):
            try:
                os.remove(temp_path)
            except OSError:
                pass


@flask_app.route("/api/export_week_pdf", methods=["POST"])
def export_week_pdf():
    if REMOTE_STORAGE:
        return jsonify({
            "status": "error",
            "message": "Экспорт PDF временно отключён на тестовом стенде удалённого хранения.",
        }), 503
    data = request.get_json() or {}
    week_start = str(data.get("week_start", "")).strip()
    if not week_start:
        today = datetime.date.today()
        start_date = today - datetime.timedelta(days=today.weekday())
    else:
        try:
            start_date = datetime.datetime.strptime(week_start, "%Y-%m-%d").date()
        except ValueError:
            return jsonify({"status": "error", "message": "Некорректная дата начала недели."}), 400

    schedule = load_json(DATA_FILE)
    week_data = {}
    for i in range(7):
        key = (start_date + datetime.timedelta(days=i)).strftime("%Y-%m-%d")
        week_data[key] = sorted(schedule.get(key, []), key=lambda x: x.get("time", "00:00"))

    try:
        pdf_stream = build_week_schedule_pdf(start_date, week_data)
    except Exception as exc:
        return jsonify({"status": "error", "message": f"Не удалось сформировать PDF: {exc}"}), 500

    request_user = getattr(g, "telegram_user", {}) or {}
    teacher_target = str(getattr(g, "teacher_id", "") or request_user.get("id", "")).strip()
    teacher_chat_id = numeric_telegram_chat_id(teacher_target)
    if teacher_chat_id is None:
        return jsonify({"status": "error", "message": "Не удалось определить ваш Telegram ID. Откройте приложение через бота."}), 400

    week_end = start_date + datetime.timedelta(days=6)
    filename = f"temli_raspisanie_{start_date.strftime('%Y-%m-%d')}_{week_end.strftime('%Y-%m-%d')}.pdf"
    fd, temp_path = tempfile.mkstemp(prefix="week_export_", suffix=".pdf")
    try:
        with os.fdopen(fd, "wb") as temp_file:
            temp_file.write(pdf_stream.getvalue())
        ok, error = send_receipt_from_flask(
            teacher_chat_id,
            temp_path,
            f"📄 TEMLI · Расписание недели {format_ru_date(start_date)} - {format_ru_date(week_end)}",
            filename=filename,
        )
        if not ok:
            return jsonify({"status": "error", "message": f"Не удалось отправить PDF в Telegram: {error}"}), 502
        return jsonify({"status": "ok", "message": "PDF расписания отправлен вам в Telegram.", "filename": filename})
    finally:
        if os.path.exists(temp_path):
            try:
                os.remove(temp_path)
            except OSError:
                pass


@flask_app.route("/api/health", methods=["GET"])
def health():
    return jsonify({
        "status": "ok",
        "message": "API работает",
        "storage": "remote-json-test" if REMOTE_STORAGE else "local",
    })


@flask_app.route("/api/get_week_schedule", methods=["POST"])
def get_week_schedule():
    data = request.get_json() or {}
    week_start = data.get("week_start")
    if not week_start:
        today = datetime.date.today()
        monday = today - datetime.timedelta(days=today.weekday())
        week_start = monday.strftime("%Y-%m-%d")

    try:
        start_date = datetime.datetime.strptime(week_start, "%Y-%m-%d").date()
    except ValueError:
        return jsonify({"status": "error", "message": "Некорректная дата начала недели."}), 400

    schedule = load_json(DATA_FILE)
    week_data = {}
    for i in range(7):
        key = (start_date + datetime.timedelta(days=i)).strftime("%Y-%m-%d")
        week_data[key] = sorted(schedule.get(key, []), key=lambda x: x.get("time", "00:00"))
    return jsonify({"status": "ok", "schedule": week_data})


@flask_app.route("/api/get_students", methods=["GET"])
def get_students():
    students = load_json(STUDENTS_FILE)
    changed = False
    teacher_id = str(getattr(g, "teacher_id", "") or "").strip()
    if teacher_id and teacher_id in students:
        # Аккаунт преподавателя не является учеником своего tenant.
        students.pop(teacher_id, None)
        save_json(STUDENTS_FILE, students)
        changed = False

    # Мягкая миграция старых данных: раньше все новые ученики получали color=0.
    # Если весь список всё ещё одноцветный, распределяем палитру один раз.
    dict_students = [(sid, raw) for sid, raw in students.items() if isinstance(raw, dict)]
    if len(dict_students) > 1:
        colors = [raw.get("color", 0) for _sid, raw in dict_students]
        if all(str(c) == "0" for c in colors):
            for idx, (_sid, raw) in enumerate(sorted(dict_students, key=lambda pair: str(pair[1].get("name", pair[0])).lower())):
                raw["color"] = idx % 7
            changed = True

    for student_id, raw in list(students.items()):
        if isinstance(raw, dict) and "contacts" not in raw:
            contacts = normalize_contacts(raw)
            if contacts:
                raw["contacts"] = contacts
                changed = True
    if changed:
        save_json(STUDENTS_FILE, students)
    return jsonify({"status": "ok", "students": students})


@flask_app.route("/api/get_settings", methods=["GET"])
def get_settings():
    settings = load_settings()
    students = load_json(STUDENTS_FILE)
    schedule = load_json(DATA_FILE)
    has_students = any(str(student_id) != current_teacher_id() for student_id in students)
    has_lessons = any(bool(lessons) for lessons in schedule.values())
    onboarding_needed = not settings.get("onboarding_completed") and not has_students and not has_lessons
    return jsonify({"status": "ok", "settings": settings, "onboarding_needed": onboarding_needed})


@flask_app.route("/api/update_settings", methods=["POST"])
def update_settings():
    data = request.get_json() or {}
    if not isinstance(data, dict):
        return jsonify({"status": "error", "message": "Ожидается объект настроек."}), 400
    return update_settings_data(data)


@serialized_data
def update_settings_data(data):
    settings = load_settings()
    boolean_keys = {"default_reminders_enabled", "default_student_reminders", "default_send_receipts", "default_send_receipt_copy", "onboarding_completed", "parent_lesson_end"}
    notification_template_keys = {
        "student_binding_template", "parent_binding_template",
        "student_reminder_template", "parent_lesson_end_template",
        "teacher_delay_template", "student_delay_template", "parent_delay_template"
    }
    text_keys = {
        "company_name", "inn", "ogrnip", "address", "phone", "service_name",
        "tax_system", "email_sender", "thanks_text", "website", "bank_name",
        "bik", "account_number", "corr_account", "recipient", "payment_comment", "zoom_link",
        "work_start", "work_end"
    } | notification_template_keys
    if any(key in data and (not isinstance(data[key], str) or len(data[key]) > 1000)
           for key in notification_template_keys):
        return jsonify({"status": "error", "message": "Текст сообщения должен быть не длиннее 1000 символов."}), 400
    enum_values = {
        "language": {"ru", "en"},
        "currency": {"RUB", "USD", "EUR", "CNY", "TRY"},
    }
    if "days_off" in data:
        raw_days_off = data.get("days_off")
        if not isinstance(raw_days_off, list):
            return jsonify({"status": "error", "message": "Выходные дни должны быть списком."}), 400
        if any(isinstance(day, bool) for day in raw_days_off):
            return jsonify({"status": "error", "message": "Некорректный день недели."}), 400
        try:
            days_off = sorted({int(day) for day in raw_days_off})
        except (TypeError, ValueError):
            return jsonify({"status": "error", "message": "Некорректный день недели."}), 400
        if any(day < 0 or day > 6 for day in days_off):
            return jsonify({"status": "error", "message": "Некорректный день недели."}), 400
        settings["days_off"] = days_off

    next_work_start = str(data.get("work_start", settings.get("work_start", "06:00"))).strip()
    next_work_end = str(data.get("work_end", settings.get("work_end", "00:00"))).strip()
    try:
        datetime.datetime.strptime(next_work_start, "%H:%M")
        datetime.datetime.strptime(next_work_end, "%H:%M")
    except ValueError:
        return jsonify({"status": "error", "message": "Некорректные рабочие часы."}), 400
    if parse_work_end_minutes(next_work_end) <= parse_minutes(next_work_start):
        return jsonify({"status": "error", "message": "Время окончания должно быть позже начала."}), 400
    for key in boolean_keys:
        if key in data:
            settings[key] = bool(data.get(key))
    for key in text_keys:
        if key in data:
            settings[key] = str(data.get(key, "")).strip()
    for key, allowed in enum_values.items():
        if key in data:
            value = str(data.get(key, "")).strip()
            if value not in allowed:
                return jsonify({"status": "error", "message": "Некорректная настройка языка или валюты."}), 400
            settings[key] = value
    save_json(SETTINGS_FILE, settings)
    return jsonify({"status": "ok", "settings": settings})


@flask_app.route("/api/upload_receipt_asset", methods=["POST"])
def upload_receipt_asset():
    if REMOTE_STORAGE:
        return jsonify({
            "status": "error",
            "message": "Загрузка файлов временно отключена на тестовом стенде удалённого хранения.",
        }), 503
    asset_type = str(request.form.get("asset_type", "")).strip()
    mapping = {
        "logo": "receipt_logo",
        "signature": "receipt_signature",
        "qrcode": "receipt_qrcode",
    }
    if asset_type not in mapping:
        return jsonify({"status": "error", "message": "Неизвестный тип файла."}), 400

    uploaded = request.files.get("file")
    if not uploaded or not uploaded.filename:
        return jsonify({"status": "error", "message": "Файл не выбран."}), 400

    raw = uploaded.read(5 * 1024 * 1024 + 1)
    if len(raw) > 5 * 1024 * 1024:
        return jsonify({"status": "error", "message": "Файл больше 5 МБ."}), 400

    ext = os.path.splitext(uploaded.filename)[1].lower()
    if ext not in {".png", ".jpg", ".jpeg"}:
        return jsonify({"status": "error", "message": "Разрешены PNG, JPG и JPEG."}), 400

    return save_receipt_asset(asset_type, ext, raw, mapping[asset_type])


@serialized_data
def save_receipt_asset(asset_type, ext, raw, setting_key):
    if REMOTE_STORAGE:
        raise RuntimeError("Бинарное хранилище тестового стенда ещё не подключено.")
    # Request body is already read; a slow upload must not hold the data lock.
    os.makedirs(current_receipt_assets_dir(), exist_ok=True)
    filename = f"{asset_type}{ext}"
    destination = os.path.join(current_receipt_assets_dir(), filename)
    fd, temporary = tempfile.mkstemp(prefix="asset_", suffix=".tmp", dir=current_receipt_assets_dir())
    try:
        with os.fdopen(fd, "wb") as f:
            f.write(raw)
            f.flush()
            os.fsync(f.fileno())
        os.replace(temporary, destination)
    finally:
        if os.path.exists(temporary):
            os.remove(temporary)

    settings = load_settings()
    settings[setting_key] = filename
    save_json(SETTINGS_FILE, settings)
    for old_ext in (".png", ".jpg", ".jpeg"):
        old_path = os.path.join(current_receipt_assets_dir(), f"{asset_type}{old_ext}")
        if old_path != destination and os.path.exists(old_path):
            os.remove(old_path)
    return jsonify({"status": "ok", "settings": settings, "filename": filename})




@flask_app.route("/api/update_student_profile", methods=["POST"])
def update_student_profile():
    data = request.get_json() or {}
    student_id = str(data.get("student_id", "")).strip()
    if not student_id:
        return jsonify({"status": "error", "message": "Не указан ученик."}), 400
    calendar_name = str(data.get("calendar_name", "")).strip()[:24]
    birthday = str(data.get("birthday", "")).strip()
    if birthday:
        try:
            datetime.datetime.strptime(birthday, "%Y-%m-%d")
        except ValueError:
            return jsonify({"status": "error", "message": "Некорректная дата рождения."}), 400
    default_price_raw = data.get("default_price", None)
    if default_price_raw is None:
        default_price = None
    else:
        try:
            default_price = max(0.0, float(default_price_raw or 0))
        except (TypeError, ValueError):
            return jsonify({"status": "error", "message": "Некорректная стоимость урока."}), 400
    status = str(data.get("status", "active")).strip().lower()
    if status not in {"active", "paused"}:
        status = "active"
    note = str(data.get("note", "")).strip()[:1000]
    board_link = str(data.get("board_link", "")).strip()[:1000]
    zoom_link = str(data.get("zoom_link", "")).strip()[:1000]
    contacts = normalize_contacts(data)
    student_contacts = normalize_student_contacts(data)
    with DATA_LOCK:
        students = load_json(STUDENTS_FILE)
        if student_id not in students:
            return jsonify({"status": "error", "message": "Ученик не найден."}), 404
        if isinstance(students[student_id], str):
            students[student_id] = {"name": students[student_id], "user_id": student_id, "color": 0}
        student = students[student_id]
        student["calendar_name"] = calendar_name
        if type(data.get('student_reminders')) is bool:
            student['student_reminders'] = data['student_reminders']
        if 'parent_lesson_end' in data:
            if type(data['parent_lesson_end']) is bool:
                student['parent_lesson_end'] = data['parent_lesson_end']
            elif data['parent_lesson_end'] is None:
                student.pop('parent_lesson_end', None)
        student["birthday"] = birthday
        if default_price is not None:
            student["default_price"] = default_price
        student["status"] = status
        student["note"] = note
        student["board_link"] = board_link
        student["zoom_link"] = zoom_link
        student["contacts"] = contacts
        student["student_contacts"] = student_contacts
        if 'parent_name' in data:
            student['parent_name'] = str(data['parent_name'] or '').strip()[:128]
        save_json(STUDENTS_FILE, students)
    return jsonify({"status": "ok", "student": student})


@flask_app.route("/api/set_student_archived", methods=["POST"])
def set_student_archived():
    data = request.get_json() or {}
    student_id = str(data.get("student_id", "")).strip()
    archived = data.get("archived")
    if not student_id or not isinstance(archived, bool):
        return jsonify({"status": "error", "message": "Укажите ученика и действие с архивом."}), 400
    with DATA_LOCK:
        students = load_json(STUDENTS_FILE)
        if student_id not in students:
            return jsonify({"status": "error", "message": "Ученик не найден."}), 404
        student = students[student_id]
        if isinstance(student, str):
            student = {"name": student}
            students[student_id] = student
        student["archived"] = archived
        save_json(STUDENTS_FILE, students)
    return jsonify({"status": "ok", "student": student})


@flask_app.route("/api/delete_student", methods=["POST"])
def delete_student():
    data = request.get_json() or {}
    student_id = str(data.get("student_id", "")).strip()
    confirm_name = str(data.get("confirm_name", "")).strip()
    if not student_id:
        return jsonify({"status": "error", "message": "Не указан ученик."}), 400
    with DATA_LOCK:
        students = load_json(STUDENTS_FILE)
        if student_id not in students:
            return jsonify({"status": "error", "message": "Ученик не найден."}), 404
        raw_student = students[student_id]
        student_name = str(raw_student if isinstance(raw_student, str) else raw_student.get("name", "")).strip()
        if not student_name or confirm_name != student_name:
            return jsonify({"status": "error", "message": "Имя для подтверждения не совпадает."}), 400

        schedule = load_json(DATA_FILE)
        old_schedule = copy.deepcopy(schedule)
        old_students = copy.deepcopy(students)
        now = receipt_now().replace(tzinfo=None)
        removed_lessons = 0
        removed_group_members = 0
        for date_key in list(schedule):
            kept = []
            for lesson in schedule[date_key]:
                try:
                    lesson_start = datetime.datetime.strptime(
                        f"{date_key} {lesson.get('time', '00:00')}", "%Y-%m-%d %H:%M")
                except (TypeError, ValueError):
                    lesson_start = datetime.datetime.min
                if lesson_start < now:
                    kept.append(lesson)
                    continue
                if lesson.get("lesson_type") == "group":
                    members = lesson.get("group_members") or []
                    remaining = [member for member in members if str(member.get("student_id", "")) != student_id]
                    if any(has_payment_link(member) for member in members
                           if str(member.get("student_id", "")) == student_id):
                        return jsonify({"status": "error", "message": "У ученика есть оплаченные будущие занятия. Сначала отмените оплату или уберите ученика в архив."}), 409
                    if len(remaining) == len(members):
                        kept.append(lesson)
                    elif remaining:
                        lesson["group_members"] = remaining
                        lesson["paid"] = all(bool(member.get("paid")) for member in remaining)
                        removed_group_members += len(members) - len(remaining)
                        kept.append(lesson)
                    else:
                        removed_lessons += 1
                elif str(lesson.get("student_id", "")) == student_id:
                    if has_payment_link(lesson):
                        return jsonify({"status": "error", "message": "У ученика есть оплаченные будущие занятия. Сначала отмените оплату или уберите ученика в архив."}), 409
                    removed_lessons += 1
                else:
                    kept.append(lesson)
            if kept:
                schedule[date_key] = kept
            else:
                schedule.pop(date_key, None)

        students.pop(student_id)
        try:
            save_json(DATA_FILE, schedule)
            save_json(STUDENTS_FILE, students)
        except Exception as exc:
            _save_json_raw(tenant_file(DATA_FILE), old_schedule)
            _save_json_raw(tenant_file(STUDENTS_FILE), old_students)
            return jsonify({"status": "error", "message": f"Удаление не сохранено, изменения восстановлены: {exc}"}), 500
    return jsonify({"status": "ok", "removed_lessons": removed_lessons,
                    "removed_group_members": removed_group_members})


@flask_app.route("/api/get_student_stats", methods=["POST"])
def get_student_stats():
    data = request.get_json() or {}
    student_id = str(data.get("student_id", "")).strip()
    if not student_id:
        return jsonify({"status": "error", "message": "Не указан ученик."}), 400
    return jsonify({"status": "ok", **student_lesson_stats(student_id)})


@flask_app.route("/api/update_lesson_report", methods=["POST"])
def update_lesson_report():
    data = request.get_json() or {}
    date = str(data.get("date", "")).strip()
    lesson_id = str(data.get("id", "")).strip()
    report = str(data.get("report", "") or "").strip()[:1000]
    if not date or not lesson_id:
        return jsonify({"status": "error", "message": "Не указаны дата или занятие."}), 400
    with DATA_LOCK:
        schedule = load_json(DATA_FILE)
        for lesson in schedule.get(date, []):
            if str(lesson.get("id", "")) == lesson_id:
                if is_personal_event(lesson):
                    return jsonify({"status": "error", "message": "У личного дела нет отчёта занятия."}), 409
                lesson["report"] = report
                save_json(DATA_FILE, schedule)
                return jsonify({"status": "ok", "report": report})
    return jsonify({"status": "error", "message": "Занятие не найдено."}), 404


@flask_app.route("/api/get_work_center", methods=["POST"])
def get_work_center():
    data = request.get_json() or {}
    week_start = data.get("week_start")
    return jsonify({"status": "ok", **build_work_center(week_start)})


@flask_app.route("/api/update_student_color", methods=["POST"])
def update_student_color():
    data = request.get_json() or {}
    student_id = str(data.get("student_id", "")).strip()
    color = data.get("color")
    if not student_id:
        return jsonify({"status": "error", "message": "Не указан ученик."}), 400

    if isinstance(color, str) and HEX_COLOR_RE.fullmatch(color):
        normalized_color = color.lower()
    else:
        try:
            normalized_color = int(color)
        except (TypeError, ValueError):
            return jsonify({"status": "error", "message": "Некорректный цвет."}), 400
        if normalized_color < 0 or normalized_color > 6:
            return jsonify({"status": "error", "message": "Некорректный цвет."}), 400

    with DATA_LOCK:
        students = load_json(STUDENTS_FILE)
        if student_id not in students:
            return jsonify({"status": "error", "message": "Ученик не найден"}), 404
        if isinstance(students[student_id], str):
            students[student_id] = {"name": students[student_id], "color": normalized_color}
        else:
            students[student_id]["color"] = normalized_color
        save_json(STUDENTS_FILE, students)
    return jsonify({"status": "ok"})


@flask_app.route("/api/add_lesson", methods=["POST"])
def add_lesson():
    data = request.get_json() or {}
    date = data.get("date")
    lesson_time = data.get("time")
    entry_type = "personal" if data.get("entry_type") == "personal" else "lesson"
    lesson_type = "group" if data.get("lesson_type") == "group" else "student"
    student = str(data.get("student", "")).strip()
    group_name = str(data.get("group_name", "")).strip()
    event_title = str(data.get("title", "")).strip()
    event_notes = str(data.get("notes", "") or "").strip()
    repeat = str(data.get("repeat", "no") or "no").strip()
    contacts = normalize_contacts(data)
    student_contacts = normalize_student_contacts(data)
    reminder_enabled = bool(data.get("reminder_enabled", True))
    raw_lesson_price = data.get("price")
    lesson_price = normalize_amount(raw_lesson_price)
    repeat_until = str(data.get("repeat_until", "") or "").strip()

    try:
        duration = int(data.get("duration", 60))
        reminder_minutes = int(data.get("reminder_minutes", 60))
        start_dt = datetime.datetime.strptime(date, "%Y-%m-%d")
        datetime.datetime.strptime(lesson_time, "%H:%M")
    except (TypeError, ValueError):
        return jsonify({"status": "error", "message": "Некорректная дата, время или числовое поле."}), 400

    if not date or not lesson_time:
        return jsonify({"status": "error", "message": "Заполните все обязательные поля."}), 400

    if repeat not in {"no", "month", "year"}:
        return jsonify({"status": "error", "message": "Некорректный режим повтора."}), 400
    if duration < 15 or duration > 1440:
        return jsonify({"status": "error", "message": "Длительность занятия должна быть от 15 до 1440 минут."}), 400
    if reminder_minutes < 0 or reminder_minutes > 10080:
        return jsonify({"status": "error", "message": "Напоминание должно быть в диапазоне от 0 до 10080 минут."}), 400
    if entry_type == "personal" and not event_title:
        return jsonify({"status": "error", "message": "Укажите название личного дела."}), 400
    if len(event_title) > 100 or len(event_notes) > 1000:
        return jsonify({"status": "error", "message": "Название или заметки слишком длинные."}), 400
    if entry_type != "personal" and raw_lesson_price not in (None, "") and lesson_price is None:
        return jsonify({"status": "error", "message": "Стоимость занятия должна быть числом."}), 400
    if entry_type != "personal" and lesson_price is not None and lesson_price < 0:
        return jsonify({"status": "error", "message": "Стоимость занятия не может быть отрицательной."}), 400

    dates_to_add = [date]
    if repeat == "month":
        dates_to_add.extend(
            (start_dt + datetime.timedelta(weeks=i)).strftime("%Y-%m-%d")
            for i in range(1, 5)
        )
    elif repeat == "year":
        if repeat_until:
            try:
                end_dt = datetime.datetime.strptime(repeat_until, "%Y-%m-%d")
            except ValueError:
                return jsonify({"status": "error", "message": "Некорректная дата окончания повторов."}), 400
        else:
            end_dt = datetime.datetime(start_dt.year if start_dt.month <= 5 else start_dt.year + 1, 5, 31)
        if end_dt < start_dt:
            return jsonify({"status": "error", "message": "Дата окончания повторов не может быть раньше первого занятия."}), 400
        if (end_dt - start_dt).days > 370:
            return jsonify({"status": "error", "message": "Серию можно создать максимум на один учебный год."}), 400
        curr = start_dt + datetime.timedelta(weeks=1)
        while curr <= end_dt:
            dates_to_add.append(curr.strftime("%Y-%m-%d"))
            curr += datetime.timedelta(weeks=1)

    with DATA_LOCK:
        students = load_json(STUDENTS_FILE)
        group_members = []
        student_id = str(data.get("student_id", "")).strip()

        if entry_type == "personal":
            pass
        elif lesson_type == "group":
            group_members = normalize_group_members(data, students)
            if not group_name:
                return jsonify({"status": "error", "message": "Укажите название группы."}), 400
            if not group_members:
                return jsonify({"status": "error", "message": "Добавьте хотя бы одного ученика в группу."}), 400
            save_json(STUDENTS_FILE, students)
        else:
            if not student:
                return jsonify({"status": "error", "message": "Укажите ученика."}), 400
            if not student_id or student_id == "manual":
                student_id = f"manual_{time.time_ns()}"
                students[student_id] = {
                    **new_student_notification_settings(),
                    "name": student,
                    "username": "",
                    "contacts": {},
                    "student_contacts": {},
                    "default_price": 0,
                    "user_id": student_id,
                    "color": next_auto_color(students),
                }
            elif student_id not in students:
                return jsonify({"status": "error", "message": "Ученик не найден."}), 404

            if isinstance(students[student_id], str):
                students[student_id] = {
                    "name": students[student_id], "contacts": {}, "student_contacts": {},
                    "default_price": 0, "color": next_auto_color(students),
                }
            student_record = students[student_id]
            if 'parent_name' in data:
                student_record['parent_name'] = str(data['parent_name'] or '').strip()[:128]
            if "contacts" in data:
                student_record["contacts"] = contacts
            if "student_contacts" in data:
                student_record["student_contacts"] = student_contacts
            save_json(STUDENTS_FILE, students)

        schedule = load_json(DATA_FILE)
        series_id = next_series_id() if len(dates_to_add) > 1 else ""
        single_price = lesson_price if entry_type == "lesson" and lesson_type == "student" and lesson_price is not None else None
        for lesson_date in dates_to_add:
            if entry_type == "personal":
                lesson = {
                    "id": next_lesson_id(), "entry_type": "personal",
                    "title": event_title, "notes": event_notes,
                    "visibility": "private", "time": lesson_time, "duration": duration,
                }
            else:
                lesson = {
                    "id": next_lesson_id(),
                    "time": lesson_time,
                    "duration": duration,
                    "lesson_type": lesson_type,
                    "price": single_price if lesson_type == "student" else "",
                    "reminder_enabled": reminder_enabled,
                    "reminder_minutes": reminder_minutes,
                    "reminder_text": data.get("reminder_text", ""),
                    "zoom_link": data.get("zoom_link", ""),
                    "paid": False,
                }
            if entry_type == "lesson" and lesson_type == "group":
                lesson["group_name"] = group_name
                lesson["student"] = group_name
                lesson["group_members"] = [dict(member, paid=False, free=False, paid_amount=0.0) for member in group_members]
            elif entry_type == "lesson":
                lesson["student"] = student
                lesson["student_id"] = student_id
            if series_id:
                lesson["series_id"] = series_id
            schedule.setdefault(lesson_date, []).append(lesson)

        save_json(DATA_FILE, schedule)

    return jsonify({"status": "ok", "students": students, "student_id": student_id if entry_type == "lesson" and lesson_type == "student" else ""})


def lesson_price_can_change(date_key, lesson, target, not_before):
    try:
        lesson_start = datetime.datetime.strptime(
            f"{date_key} {lesson.get('time', '00:00')}", "%Y-%m-%d %H:%M"
        )
        if getattr(not_before, "tzinfo", None):
            lesson_start = lesson_start.replace(tzinfo=not_before.tzinfo)
    except (TypeError, ValueError):
        return False
    if lesson_start < not_before or lesson.get("cancelled"):
        return False
    return not (
        target.get("paid")
        or target.get("free")
        or target.get("paid_via_subscription")
        or target.get("allocation_ids")
        or (normalize_amount(target.get("paid_amount")) or 0) > 0
    )


@flask_app.route("/api/update_lesson", methods=["POST"])
def update_lesson():
    data = request.get_json() or {}
    date = data.get("date")
    lesson_id = data.get("id")
    if not date or not lesson_id:
        return jsonify({"status": "error", "message": "Не указаны дата или ID занятия."}), 400
    price_scope = str(data.get("price_scope", "single") or "single").strip()
    if price_scope not in {"single", "future"}:
        return jsonify({"status": "error", "message": "Неизвестный вариант изменения цены."}), 400

    with DATA_LOCK:
        schedule = load_json(DATA_FILE)
        students = load_json(STUDENTS_FILE)
        found = False
        updated_prices = 0
        now = receipt_now()

        if date in schedule:
            for lesson in schedule[date]:
                if lesson.get("id") == lesson_id:
                    if is_personal_event(lesson):
                        if data.get("entry_type", "personal") != "personal":
                            return jsonify({"status": "error", "message": "Тип личного дела изменить нельзя."}), 409
                        title = str(data.get("title", lesson.get("title", ""))).strip()
                        notes = str(data.get("notes", lesson.get("notes", "")) or "").strip()
                        try:
                            duration = int(data.get("duration", lesson.get("duration", 60)))
                            next_time = str(data.get("time", lesson.get("time", ""))).strip()
                            datetime.datetime.strptime(next_time, "%H:%M")
                        except (TypeError, ValueError):
                            return jsonify({"status": "error", "message": "Некорректное время или длительность."}), 400
                        if not title or len(title) > 100 or len(notes) > 1000 or duration < 15 or duration > 1440:
                            return jsonify({"status": "error", "message": "Проверьте название, заметки и длительность личного дела."}), 400
                        lesson.update(entry_type="personal", title=title, notes=notes,
                                      visibility="private", time=next_time, duration=duration)
                        save_json(DATA_FILE, schedule)
                        return jsonify({"status": "ok", "message": "Личное дело обновлено.", "students": students, "updated_prices": 0})
                    if data.get("entry_type") == "personal":
                        return jsonify({"status": "error", "message": "Тип занятия изменить нельзя."}), 409
                    lesson_type = "group" if data.get("lesson_type", lesson.get("lesson_type")) == "group" else "student"
                    old_type = "group" if lesson.get("lesson_type") == "group" else "student"
                    if old_type != lesson_type and lesson_has_payment_link(lesson):
                        return jsonify({"status": "error", "message": "Сначала отмените оплату, чтобы изменить тип занятия."}), 409
                    if old_type == lesson_type == "student" and has_payment_link(lesson):
                        if str(data.get("student_id", lesson.get("student_id", ""))) != str(lesson.get("student_id", "")):
                            return jsonify({"status": "error", "message": "Сначала отмените оплату, чтобы заменить ученика."}), 409
                    lesson["lesson_type"] = lesson_type
                    lesson["reminder_enabled"] = bool(data.get("reminder_enabled", lesson.get("reminder_enabled", True)))
                    lesson["reminder_minutes"] = max(0, int(data.get("reminder_minutes", lesson.get("reminder_minutes", 60))))
                    lesson["reminder_text"] = data.get("reminder_text", lesson.get("reminder_text", ""))
                    lesson["zoom_link"] = data.get("zoom_link", lesson.get("zoom_link", ""))
                    lesson.pop("reminder_sent_for", None)
                    future_prices = {}
                    if "price" in data and lesson_type != "group":
                        parsed_price = normalize_amount(data.get("price"))
                        new_price = parsed_price if parsed_price is not None and parsed_price >= 0 else ""
                        if new_price != lesson.get("price"):
                            if not lesson_price_can_change(date, lesson, lesson, now):
                                return jsonify({"status": "error", "message": "Нельзя изменить цену прошедшего, отменённого, бесплатного или оплаченного занятия."}), 409
                            lesson["price"] = new_price
                            updated_prices += 1
                        if price_scope == "future" and lesson.get("student_id"):
                            future_prices[str(lesson.get("student_id"))] = new_price

                    if lesson_type == "group":
                        group_name = str(data.get("group_name", lesson.get("group_name", ""))).strip()
                        members = normalize_group_members(data, students)
                        if not group_name or not members:
                            return jsonify({"status": "error", "message": "Для группы нужны название и хотя бы один ученик."}), 400
                        old_by_id = {str(m.get("student_id")): m for m in lesson.get("group_members", []) if isinstance(m, dict)}
                        new_ids = {str(m.get("student_id")) for m in members}
                        if any(has_payment_link(old) for sid, old in old_by_id.items() if sid not in new_ids):
                            return jsonify({"status": "error", "message": "Сначала отмените оплату участника, чтобы удалить его из группы."}), 409
                        for member in members:
                            old = old_by_id.get(str(member.get("student_id")), {})
                            submitted_price = member.get("price")
                            member["paid"] = bool(old.get("paid", False))
                            member["free"] = bool(old.get("free", False))
                            member["paid_amount"] = max(0.0, normalize_amount(old.get("paid_amount")) or 0.0)
                            # Цена из формы редактирования имеет приоритет; для старых участников
                            # без переданной цены сохраняется старая.
                            if member.get("price") in ("", None) and old.get("price") not in ("", None):
                                member["price"] = old.get("price")
                            for key in ("receipt_number", "receipt_path", "receipt_created_at", "receipt_logged", "paid_via_subscription", "allocation_ids"):
                                if old.get(key) is not None:
                                    member[key] = old.get(key)
                            if submitted_price != old.get("price"):
                                if not lesson_price_can_change(date, lesson, old, now):
                                    return jsonify({"status": "error", "message": f"Нельзя изменить цену у {member.get('name', 'ученика')}: занятие уже прошло, отменено или оплачено."}), 409
                                updated_prices += 1
                                if price_scope == "future" and member.get("student_id"):
                                    future_prices[str(member.get("student_id"))] = member.get("price")
                        lesson["group_name"] = group_name
                        lesson["student"] = group_name
                        lesson["group_members"] = members
                        lesson["paid"] = bool(members) and all(bool(m.get("paid")) for m in members)
                        lesson.pop("student_id", None)
                        lesson.pop("contacts", None)
                        save_json(STUDENTS_FILE, students)
                    else:
                        lesson["student"] = data.get("student", lesson.get("student", ""))
                        lesson["student_id"] = data.get("student_id", lesson.get("student_id", ""))
                        # Financial status is changed only through payment endpoints.
                        lesson.pop("group_name", None)
                        lesson.pop("group_members", None)
                    if price_scope == "future" and future_prices:
                        try:
                            selected_start = datetime.datetime.strptime(
                                f"{date} {lesson.get('time', '00:00')}", "%Y-%m-%d %H:%M"
                            ).replace(tzinfo=now.tzinfo)
                        except (TypeError, ValueError):
                            selected_start = now
                        cutoff = max(now, selected_start)
                        for future_date, future_lessons in schedule.items():
                            for future_lesson in future_lessons:
                                if future_lesson.get("lesson_type") == "group":
                                    for future_member in future_lesson.get("group_members", []):
                                        sid = str(future_member.get("student_id", ""))
                                        if sid in future_prices and lesson_price_can_change(future_date, future_lesson, future_member, cutoff):
                                            new_price = future_prices[sid]
                                            if future_member.get("price") != new_price:
                                                future_member["price"] = new_price
                                                updated_prices += 1
                                else:
                                    sid = str(future_lesson.get("student_id", ""))
                                    if sid in future_prices and lesson_price_can_change(future_date, future_lesson, future_lesson, cutoff):
                                        new_price = future_prices[sid]
                                        if future_lesson.get("price") != new_price:
                                            future_lesson["price"] = new_price
                                            updated_prices += 1
                    found = True
                    break

        if found:
            save_json(DATA_FILE, schedule)
            return jsonify({"status": "ok", "message": "Занятие обновлено.", "students": students, "updated_prices": updated_prices})
    return jsonify({"status": "error", "message": "Занятие не найдено."}), 404


@flask_app.route("/api/pay_subscription", methods=["POST"])
def pay_subscription():
    data = request.get_json() or {}
    date = data.get("date")
    lesson_id = data.get("id")
    send_receipt = bool(data.get("send_receipt", True))
    try:
        amount = float(data.get("amount", 0) or 0)
    except (TypeError, ValueError):
        return jsonify({"status": "error", "message": "Некорректная сумма абонемента."}), 400
    if not math.isfinite(amount) or amount <= 0:
        return jsonify({"status": "error", "message": "Сумма абонемента должна быть больше нуля."}), 400

    requested_lesson_count = data.get("lesson_count")
    if requested_lesson_count in (None, ""):
        requested_lesson_count = None
    else:
        try:
            requested_lesson_count = int(requested_lesson_count)
        except (TypeError, ValueError):
            return jsonify({"status": "error", "message": "Некорректное количество занятий."}), 400
        if requested_lesson_count < 2:
            return jsonify({"status": "error", "message": "Для абонемента укажите минимум 2 занятия."}), 400

    with DATA_LOCK:
        schedule = load_json(DATA_FILE)
        students = load_json(STUDENTS_FILE)
        lesson = next((item for item in schedule.get(date, []) if item.get("id") == lesson_id), None)
        if not lesson:
            return jsonify({"status": "error", "message": "Занятие не найдено."}), 404
        if is_personal_event(lesson):
            return jsonify({"status": "error", "message": "Оплата для личного дела недоступна."}), 409
        requested_student_id = str(data.get("student_id", "") or "").strip()
        if lesson.get("lesson_type") == "group":
            student_id = requested_student_id
            member = next((m for m in (lesson.get("group_members") or []) if str(m.get("student_id", "")) == student_id), None)
            if not member:
                return jsonify({"status": "error", "message": "Выберите участника группы для абонемента."}), 400
            price = member_lesson_price(students, member, lesson)
            client_name = str(member.get("name") or get_student_record(students, student_id).get("name", "Клиент"))
        else:
            student_id = str(lesson.get("student_id", ""))
            price = get_student_price(students, student_id, lesson.get("price"))
            client_name = str(lesson.get("student") or get_student_record(students, student_id).get("name", "Клиент"))
        source_is_group = lesson.get("lesson_type") == "group"
        source_target = member if source_is_group else lesson
        if source_target.get("paid") or source_target.get("free") or has_allocation(source_target) or source_target.get("paid_via_subscription"):
            return jsonify({"status": "error", "message": "Сначала отмените существующую оплату или бесплатный статус занятия."}), 409
        if price is None or price <= 0:
            return jsonify({"status": "error", "message": "Сначала укажите стоимость этого занятия."}), 400
        if requested_lesson_count is not None:
            lesson_count = requested_lesson_count
        else:
            count_exact = amount / price
            lesson_count = int(round(count_exact))
            if lesson_count < 2:
                return jsonify({"status": "error", "message": "Для абонемента сумма должна быть как минимум за 2 занятия."}), 400
            if abs(count_exact - lesson_count) > 0.0001:
                return jsonify({"status": "error", "message": f"Сумма должна делиться на стоимость урока {price:.0f} ₽ без остатка или укажите количество занятий вручную."}), 400

        # An abonnement follows the same ledger rule as a common payment:
        # oldest chargeable debt first, across individual and group lessons.
        # The lesson where the action was opened only identifies the student.
        candidates = []
        for date_key, item, target in student_payment_lessons(schedule, student_id):
            if item.get("cancelled") or target.get("paid") or target.get("free") or target.get("paid_via_subscription"):
                continue
            try:
                item_dt = datetime.datetime.strptime(f"{date_key} {item.get('time', '00:00')}", "%Y-%m-%d %H:%M")
            except (TypeError, ValueError):
                continue
            candidates.append((item_dt, date_key, item, target if item.get("lesson_type") == "group" else None))
        candidates.sort(key=lambda row: (row[0], str(row[2].get("id", ""))))
        if any(has_allocation(member if member is not None else item) for _dt, _date, item, member in candidates[:lesson_count]):
            return jsonify({"status": "error", "message": "В выбранных занятиях есть распределение общей суммы. Сначала отмените его в карточке ученика."}), 409
        if len(candidates) < lesson_count:
            return jsonify({"status": "error", "message": f"В расписании только {len(candidates)} неоплаченных занятий. Для суммы {amount:.0f} ₽ нужно {lesson_count}."}), 400

        settings = load_settings()
        now = receipt_now()
        receipt_number = next_receipt_number(now)
        try:
            receipt_path, receipt_number, created_at = generate_receipt_pdf(
                settings,
                client_name,
                amount,
                f"subscription_{lesson_id}_{time.time_ns()}",
                receipt_number=receipt_number,
                created_at=now,
                service_name_override=f"Абонемент — {lesson_count} занятий",
            )
        except Exception as exc:
            return jsonify({"status": "error", "message": f"Не удалось сформировать чек: {exc}"}), 500

        marked_ids = []
        try:
            with payment_files_transaction():
                add_receipt_to_book(
                    client_name,
                    amount,
                    receipt_number,
                    created_at,
                    status=f"Абонемент · {lesson_count} занятий",
                )
                for _dt, date_key, item, member in candidates[:lesson_count]:
                    if member is None:
                        item["paid"] = True
                        item["paid_via_subscription"] = receipt_number
                        marked_ids.append(item.get("id"))
                    else:
                        member["paid"] = True
                        member["paid_via_subscription"] = receipt_number
                        item["paid"] = bool(item.get("group_members")) and all(bool(m.get("paid")) for m in item.get("group_members") or [])
                        marked_ids.append(f"{item.get('id')}:{member.get('student_id')}")
                save_json(DATA_FILE, schedule)
        except Exception as exc:
            delete_receipt_file(receipt_path)
            return jsonify({"status": "error", "message": f"Абонемент не сохранён, изменения отменены: {exc}"}), 500

        send_teacher_copy = bool(settings.get("default_send_receipt_copy", True))
        request_user = getattr(g, "telegram_user", {}) or {}
        teacher_target = str(getattr(g, "teacher_id", "") or request_user.get("id", "")).strip()
        teacher_chat_id = numeric_telegram_chat_id(teacher_target) if send_teacher_copy else None
        student_info = get_student_record(students, student_id)
        parent_contacts = student_info.get("contacts") or lesson.get("contacts") or {}
        parent_chat_id = numeric_telegram_chat_id(parent_contacts.get("tg") if isinstance(parent_contacts, dict) else "") if send_receipt else None

    messages = []
    try:
        if send_receipt:
            if parent_chat_id is None:
                messages.append("Родителю чек не отправлен: нет числового Telegram ID.")
            else:
                ok, error = send_receipt_from_flask(parent_chat_id, receipt_path, f"Чек за абонемент: {lesson_count} занятий · {amount:.2f} руб. · № {receipt_number}")
                messages.append("Чек отправлен родителю." if ok else f"Родителю чек отправить не удалось: {error}")
        else:
            messages.append("Родителю чек не отправлялся.")
        if send_teacher_copy:
            if teacher_chat_id is None:
                messages.append("Копия вам не отправлена: откройте этого бота и нажмите /start.")
            else:
                ok, error = send_receipt_from_flask(teacher_chat_id, receipt_path, f"Копия чека: абонемент {lesson_count} занятий · {amount:.2f} руб. · № {receipt_number}")
                messages.append("Копия чека отправлена вам." if ok else f"Копию вам отправить не удалось: {error}")
    finally:
        delete_receipt_file(receipt_path)

    return jsonify({
        "status": "ok",
        "receipt_number": receipt_number,
        "lessons_paid": lesson_count,
        "marked_ids": marked_ids,
        "receipt_message": " ".join(messages),
        **student_lesson_stats(student_id),
    })


@flask_app.route("/api/mark_paid", methods=["POST"])
def mark_paid():
    data = request.get_json() or {}
    date = data.get("date")
    lesson_id = data.get("id")
    paid = bool(data.get("paid", True))
    send_receipt = bool(data.get("send_receipt", True))

    with DATA_LOCK:
        schedule = load_json(DATA_FILE)
        lesson = None
        if date in schedule:
            lesson = next((item for item in schedule[date] if item.get("id") == lesson_id), None)
        if not lesson:
            return jsonify({"status": "error", "message": "Занятие не найдено"}), 404
        if is_personal_event(lesson):
            return jsonify({"status": "error", "message": "Оплата для личного дела недоступна."}), 409

        settings = load_settings()
        send_teacher_copy = bool(settings.get("default_send_receipt_copy", True))
        request_user = getattr(g, "telegram_user", {}) or {}
        teacher_target = str(getattr(g, "teacher_id", "") or request_user.get("id", "")).strip()
        teacher_chat_id = numeric_telegram_chat_id(teacher_target) if send_teacher_copy else None

        if lesson.get("lesson_type") == "group":
            members = lesson.get("group_members") or []
            selected_ids = {str(x) for x in (data.get("paid_student_ids") or [])}
            selected_student = str(data.get("student_id", "") or "").strip()
            if selected_student:
                if not any(str(m.get("student_id", "")) == selected_student for m in members):
                    return jsonify({"status": "error", "message": "Участник группы не найден."}), 404
                selected_ids = {str(m.get("student_id", "")) for m in members if m.get("paid")}
                if paid:
                    selected_ids.add(selected_student)
                else:
                    selected_ids.discard(selected_student)
                paid = True
            for member in members:
                sid = str(member.get("student_id", ""))
                wants_paid = paid and sid in selected_ids
                if selected_student and sid != selected_student:
                    continue
                if member.get("free") and wants_paid:
                    return jsonify({"status": "error", "message": "Сначала отмените бесплатный статус."}), 409
                if (selected_student or not paid or wants_paid != bool(member.get("paid"))) and (has_allocation(member) or member.get("paid_via_subscription")):
                    return jsonify({"status": "error", "message": "Оплата общей суммой или абонементом: отдельное изменение запрещено. Откройте карточку ученика."}), 409
            students = load_json(STUDENTS_FILE)
            if paid:
                missing_prices = []
                for member in members:
                    sid = str(member.get("student_id", ""))
                    if sid in selected_ids and not bool(member.get("paid", False)) and not bool(member.get("free", False)):
                        member_amount = member_lesson_price(students, member, lesson)
                        if member_amount is None:
                            missing_prices.append(member.get("name", "Ученик"))
                if missing_prices:
                    return jsonify({"status": "error", "message": f"Укажите стоимость занятия для: {', '.join(missing_prices)}."}), 400

            for member in members:
                sid = str(member.get("student_id", ""))
                should_remain_paid = paid and sid in selected_ids
                if should_remain_paid or not bool(member.get("paid", False)) or bool(member.get("free", False)):
                    continue
                if member.get("paid_via_subscription"):
                    return jsonify({
                        "status": "error",
                        "message": f"Занятие участника {member.get('name', 'ученика')} оплачено абонементом. Отдельная отмена заблокирована.",
                    }), 409
                allocated_amount = max(0.0, normalize_amount(member.get("paid_amount")) or 0.0)
                if allocated_amount > 0 and not member.get("receipt_number"):
                    return jsonify({
                        "status": "error",
                        "message": f"Оплата участника {member.get('name', 'ученика')} распределена общей суммой. Для отмены нужна история платежа.",
                    }), 409

            receipt_jobs = []
            reversal_jobs = []
            created_receipts = []
            for member in members:
                sid = str(member.get("student_id", ""))
                if member.get("free"):
                    continue
                should_be_paid = sid in selected_ids if paid else False
                was_paid = bool(member.get("paid", False))
                member["paid"] = should_be_paid
                if not should_be_paid:
                    if was_paid:
                        reversal_amount = member_lesson_price(students, member, lesson) or 0.0
                        receipt_number = str(member.get("receipt_number", "") or "")
                        if member.get("receipt_logged") and receipt_number and reversal_amount > 0:
                            reversal_jobs.append({
                                "name": member.get("name", "Клиент"),
                                "amount": reversal_amount,
                                "number": receipt_number,
                            })
                        member["paid_amount"] = 0.0
                        for key in (
                            "receipt_number", "receipt_created_at", "receipt_logged",
                            "receipt_book_error", "paid_via_subscription",
                        ):
                            member.pop(key, None)
                    continue
                if was_paid:
                    continue

                amount = member_lesson_price(students, member, lesson)
                receipt_number = member.get("receipt_number")
                created_at = None
                if member.get("receipt_created_at"):
                    try:
                        created_at = datetime.datetime.fromisoformat(member.get("receipt_created_at"))
                    except ValueError:
                        created_at = None

                old_path = member.pop("receipt_path", None)
                delete_receipt_file(old_path)
                try:
                    receipt_path, receipt_number, created_at = generate_receipt_pdf(
                        settings, member.get("name", "Клиент"), amount, f"{lesson_id}_{sid}",
                        receipt_number=receipt_number, created_at=created_at,
                    )
                except Exception as exc:
                    for job in receipt_jobs:
                        delete_receipt_file(job["path"])
                    return jsonify({"status": "error", "message": f"Не удалось сформировать чек для {member.get('name', 'ученика')}: {exc}"}), 500

                member["receipt_number"] = receipt_number
                member["receipt_created_at"] = (created_at or receipt_now()).isoformat(timespec="seconds")
                created_receipts.append({"student_id": sid, "name": member.get("name", "Ученик"), "receipt_number": receipt_number})

                parent_chat_id = None
                if send_receipt:
                    raw_student = students.get(sid, {})
                    info = raw_student if isinstance(raw_student, dict) else {}
                    contacts = info.get("contacts") or {}
                    parent_chat_id = numeric_telegram_chat_id(contacts.get("tg") if isinstance(contacts, dict) else "")
                receipt_jobs.append({
                    "path": receipt_path,
                    "name": member.get("name", "Ученик"),
                    "number": receipt_number,
                    "amount": amount,
                    "created_at": created_at or receipt_now(),
                    "member": member,
                    "needs_book": not bool(member.get("receipt_logged")),
                    "parent_chat_id": parent_chat_id,
                    "parent_caption": f"Чек об оплате занятия для {member.get('name', 'ученика')} на {amount:.2f} руб. № {receipt_number}",
                    "teacher_caption": f"Копия чека: {member.get('name', 'ученик')} · {amount:.2f} руб. · № {receipt_number}",
                })

            lesson["paid"] = bool(members) and all(bool(member.get("paid")) or bool(member.get("free")) for member in members)
            try:
                with payment_files_transaction():
                    for job in reversal_jobs:
                        add_receipt_to_book(
                            job["name"],
                            -job["amount"],
                            job["number"],
                            receipt_now(),
                            status="Отмена оплаты участника группы",
                        )
                    for job in receipt_jobs:
                        if job["needs_book"]:
                            add_receipt_to_book(
                                job["name"],
                                job["amount"],
                                job["number"],
                                job["created_at"],
                            )
                            job["member"]["receipt_logged"] = True
                    save_json(DATA_FILE, schedule)
            except Exception as exc:
                for job in receipt_jobs:
                    delete_receipt_file(job["path"])
                return jsonify({"status": "error", "message": f"Оплата группы не сохранена, изменения отменены: {exc}"}), 500
        else:
            if lesson.get("free") or has_allocation(lesson) or lesson.get("paid_via_subscription"):
                return jsonify({"status": "error", "message": "Сначала отмените бесплатный статус, общую оплату или абонемент. Общая оплата отменяется в карточке ученика."}), 409
            if not paid:
                if lesson.get("paid_via_subscription"):
                    return jsonify({
                        "status": "error",
                        "message": "Это занятие оплачено абонементом. Нельзя отменить только одно занятие без перерасчёта всего абонемента.",
                    }), 409
                allocated_amount = max(0.0, normalize_amount(lesson.get("paid_amount")) or 0.0)
                if allocated_amount > 0 and not lesson.get("receipt_number"):
                    return jsonify({
                        "status": "error",
                        "message": "Оплата распределена общей суммой. Для безопасной отмены нужна история распределения платежа.",
                    }), 409

                students = load_json(STUDENTS_FILE)
                reversal_amount = get_student_price(students, lesson.get("student_id"), lesson.get("price")) or 0.0
                receipt_number = str(lesson.get("receipt_number", "") or "")
                try:
                    with payment_files_transaction():
                        if lesson.get("receipt_logged") and receipt_number and reversal_amount > 0:
                            add_receipt_to_book(
                                lesson.get("student", "Клиент"),
                                -reversal_amount,
                                receipt_number,
                                receipt_now(),
                                status="Отмена оплаты занятия",
                            )
                        lesson["paid"] = False
                        lesson["paid_amount"] = 0.0
                        for key in (
                            "receipt_number", "receipt_created_at", "receipt_logged",
                            "receipt_book_error", "paid_via_subscription",
                        ):
                            lesson.pop(key, None)
                        save_json(DATA_FILE, schedule)
                except Exception as exc:
                    return jsonify({"status": "error", "message": f"Отмена оплаты не сохранена, изменения отменены: {exc}"}), 500
                return jsonify({"status": "ok", "paid": False, "receipt_sent": False})

            students = load_json(STUDENTS_FILE)
            amount = get_student_price(students, lesson.get("student_id"), lesson.get("price"))
            if amount is None or amount <= 0:
                return jsonify({"status": "error", "message": "Сначала укажите стоимость этого занятия."}), 400

            receipt_number = lesson.get("receipt_number")
            receipt_created_at_raw = lesson.get("receipt_created_at")
            created_at = None
            if receipt_created_at_raw:
                try:
                    created_at = datetime.datetime.fromisoformat(receipt_created_at_raw)
                except ValueError:
                    created_at = None

            old_path = lesson.pop("receipt_path", None)
            delete_receipt_file(old_path)
            try:
                receipt_path, receipt_number, created_at = generate_receipt_pdf(
                    settings, lesson.get("student", "Клиент"), amount, lesson_id,
                    receipt_number=receipt_number, created_at=created_at,
                )
            except Exception as exc:
                return jsonify({"status": "error", "message": f"Не удалось сформировать чек: {exc}"}), 500

            try:
                with payment_files_transaction():
                    lesson["paid"] = True
                    lesson["receipt_number"] = receipt_number
                    lesson["receipt_created_at"] = (created_at or receipt_now()).isoformat(timespec="seconds")
                    if not lesson.get("receipt_logged"):
                        add_receipt_to_book(
                            lesson.get("student", "Клиент"),
                            amount,
                            receipt_number,
                            created_at or receipt_now(),
                        )
                        lesson["receipt_logged"] = True
                    lesson.pop("receipt_book_error", None)
                    save_json(DATA_FILE, schedule)
            except Exception as exc:
                delete_receipt_file(receipt_path)
                return jsonify({"status": "error", "message": f"Оплата не сохранена, изменения отменены: {exc}"}), 500

    if lesson.get("lesson_type") == "group":
        sent_names = []
        failed_names = []
        teacher_sent_names = []
        teacher_failed_names = []
        for job in receipt_jobs:
            try:
                if send_receipt:
                    if job["parent_chat_id"] is None:
                        failed_names.append(job["name"])
                    else:
                        ok, _error = send_receipt_from_flask(job["parent_chat_id"], job["path"], job["parent_caption"])
                        (sent_names if ok else failed_names).append(job["name"])

                if send_teacher_copy:
                    if teacher_chat_id is None:
                        teacher_failed_names.append(job["name"])
                    else:
                        ok, _error = send_receipt_from_flask(teacher_chat_id, job["path"], job["teacher_caption"])
                        (teacher_sent_names if ok else teacher_failed_names).append(job["name"])
            finally:
                delete_receipt_file(job["path"])

        parts = []
        if send_receipt:
            if sent_names:
                parts.append(f"Родителям отправлено: {', '.join(sent_names)}.")
            if failed_names:
                parts.append(f"Не удалось отправить родителям: {', '.join(failed_names)}.")
        else:
            parts.append("Родителям чеки не отправлялись.")
        if send_teacher_copy:
            if teacher_sent_names:
                parts.append(f"Копии вам отправлены: {len(teacher_sent_names)}.")
            elif receipt_jobs:
                parts.append("Копии вам не отправлены: откройте этого бота и нажмите /start.")

        return jsonify({
            "status": "ok",
            "paid": lesson.get("paid", False),
            "group": True,
            "members": lesson.get("group_members", []),
            "receipts_created": created_receipts,
            "receipt_sent_names": sent_names,
            "receipt_failed_names": failed_names,
            "teacher_receipt_sent_names": teacher_sent_names,
            "teacher_receipt_failed_names": teacher_failed_names,
            "receipt_message": " ".join(parts),
        })

    receipt_sent = False
    teacher_receipt_sent = False
    messages = []
    try:
        if send_receipt:
            student_info = get_student_record(load_json(STUDENTS_FILE), lesson.get("student_id"))
            contacts = student_info.get("contacts") or lesson.get("contacts") or {}
            chat_id = numeric_telegram_chat_id(contacts.get("tg") if isinstance(contacts, dict) else "")
            if chat_id is None:
                messages.append("Родителю чек не отправлен: нет числового Telegram ID.")
            else:
                caption = f"Чек об оплате занятия с {lesson.get('student', 'учеником')} на {amount:.2f} руб. № {receipt_number}"
                receipt_sent, error = send_receipt_from_flask(chat_id, receipt_path, caption)
                messages.append("Чек отправлен родителю." if receipt_sent else f"Родителю чек отправить не удалось: {error}")
        else:
            messages.append("Родителю чек не отправлялся.")

        if send_teacher_copy:
            if teacher_chat_id is None:
                messages.append("Копия вам не отправлена: откройте этого бота и нажмите /start.")
            else:
                teacher_caption = f"Копия чека: {lesson.get('student', 'ученик')} · {amount:.2f} руб. · № {receipt_number}"
                teacher_receipt_sent, error = send_receipt_from_flask(teacher_chat_id, receipt_path, teacher_caption)
                messages.append("Копия чека отправлена вам." if teacher_receipt_sent else f"Копию вам отправить не удалось: {error}")
    finally:
        delete_receipt_file(receipt_path)

    return jsonify({
        "status": "ok",
        "paid": True,
        "receipt_created": True,
        "receipt_number": receipt_number,
        "receipt_sent": receipt_sent,
        "teacher_receipt_sent": teacher_receipt_sent,
        "receipt_message": " ".join(messages),
    })



@flask_app.route("/api/set_lesson_state", methods=["POST"])
def set_lesson_state():
    data = request.get_json() or {}
    date = str(data.get("date", "") or "").strip()
    lesson_id = str(data.get("id", "") or "").strip()
    action = str(data.get("action", "") or "").strip()
    student_id = str(data.get("student_id", "") or "").strip()
    if action not in {"free", "unfree", "cancel", "restore"}:
        return jsonify({"status": "error", "message": "Неизвестное действие."}), 400
    with DATA_LOCK:
        schedule = load_json(DATA_FILE)
        lesson = next((x for x in schedule.get(date, []) if str(x.get("id", "")) == lesson_id), None)
        if not lesson:
            return jsonify({"status": "error", "message": "Занятие не найдено."}), 404
        if is_personal_event(lesson):
            return jsonify({"status": "error", "message": "Статусы занятия для личного дела недоступны."}), 409
        if action in {"cancel", "restore"}:
            lesson["cancelled"] = action == "cancel"
            if action == "cancel":
                lesson.pop("reminder_sent_for", None)
            save_json(DATA_FILE, schedule)
            return jsonify({"status": "ok", "cancelled": bool(lesson.get("cancelled"))})
        if lesson.get("lesson_type") == "group":
            member = next((m for m in (lesson.get("group_members") or []) if str(m.get("student_id", "")) == student_id), None)
            if not member:
                return jsonify({"status": "error", "message": "Участник группы не найден."}), 404
            if has_allocation(member) or member.get("paid_via_subscription") or (member.get("paid") and not member.get("free")):
                return jsonify({"status": "error", "message": "Сначала отмените оплату занятия."}), 409
            is_free = action == "free"
            member["free"] = is_free
            member["paid"] = is_free or bool(member.get("paid", False))
            if is_free:
                member["paid_amount"] = 0.0
            else:
                member["paid"] = False
            lesson["paid"] = bool(lesson.get("group_members")) and all(bool(m.get("paid")) or bool(m.get("free")) for m in lesson.get("group_members") or [])
        else:
            is_free = action == "free"
            if has_allocation(lesson) or lesson.get("paid_via_subscription") or (lesson.get("paid") and not lesson.get("free")):
                return jsonify({"status": "error", "message": "Сначала отмените оплату занятия."}), 409
            lesson["free"] = is_free
            lesson["paid"] = is_free or bool(lesson.get("paid", False))
            if is_free:
                lesson["paid_amount"] = 0.0
            else:
                lesson["paid"] = False
        save_json(DATA_FILE, schedule)
    return jsonify({"status": "ok"})


FINANCIAL_KEYS = ("paid", "free", "paid_amount", "paid_via_subscription", "allocation_ids",
                  "receipt_number", "receipt_created_at", "receipt_logged")


def has_allocation(target):
    return bool(target.get("allocation_ids")) or (normalize_amount(target.get("paid_amount")) or 0) > 0


def has_payment_link(target):
    return bool(has_allocation(target) or target.get("paid_via_subscription")
                or target.get("receipt_logged") or target.get("receipt_number")
                or (target.get("paid") and not target.get("free")))


def lesson_has_payment_link(lesson):
    if is_personal_event(lesson):
        name = str(lesson.get("title") or "Личное дело")
    elif lesson.get("lesson_type") == "group":
        return any(has_payment_link(member) for member in lesson.get("group_members", []))
    return has_payment_link(lesson)


def financial_snapshot(target):
    return {key: target[key] for key in FINANCIAL_KEYS if key in target}


def comparable_finances(target):
    return {**{"paid": False, "free": False, "paid_amount": 0.0}, **financial_snapshot(target)}


def student_payment_lessons(schedule, student_id):
    for date, lessons in schedule.items():
        for lesson in lessons:
            if lesson.get("lesson_type") == "group":
                for member in lesson.get("group_members") or []:
                    if str(member.get("student_id", "")) == student_id:
                        yield date, lesson, member
            elif str(lesson.get("student_id", "")) == student_id:
                yield date, lesson, lesson


@flask_app.route("/api/get_student_payments", methods=["POST"])
def get_student_payments():
    student_id = str((request.get_json() or {}).get("student_id", "")).strip()
    if not student_id:
        return jsonify({"status": "error", "message": "Не указан ученик."}), 400
    with DATA_LOCK:
        recover_payment_transaction()
        schedule = load_json(DATA_FILE)
        students = load_json(STUDENTS_FILE)
        history = load_json(payments_file(), {})
        lessons = []
        for date, lesson, target in student_payment_lessons(schedule, student_id):
            is_group = lesson.get("lesson_type") == "group"
            price = member_lesson_price(students, target, lesson) if is_group else get_student_price(students, student_id, lesson.get("price"))
            source = ("free" if target.get("free") else "allocation" if has_allocation(target)
                      else "subscription" if target.get("paid_via_subscription") else "direct" if target.get("paid") else "unpaid")
            lessons.append({"date": date, "id": lesson.get("id"), "time": lesson.get("time"),
                            "lesson_type": "group" if is_group else "student", "student_id": student_id,
                            "student": get_student_record(students, student_id).get("name", "Ученик"),
                            "group_name": lesson.get("group_name", ""), "price": price,
                            "cancelled": bool(lesson.get("cancelled")), "source": source,
                            **financial_snapshot(target)})
        lessons.sort(key=lambda row: (row["date"], row.get("time") or "", str(row["id"])))
        transactions = [tx for tx in history.values()
                        if tx.get("student_id") == student_id and not tx.get("request_only")]
        transactions.sort(key=lambda tx: tx["created_at"], reverse=True)
        return jsonify({"status": "ok", "lessons": lessons, "transactions": transactions})


@flask_app.route("/api/reverse_student_payment", methods=["POST"])
def reverse_student_payment():
    data = request.get_json() or {}
    student_id = str(data.get("student_id", "")).strip()
    transaction_id = str(data.get("transaction_id", ""))
    with DATA_LOCK:
        recover_payment_transaction()
        history = load_json(payments_file(), {})
        tx = history.get(transaction_id)
        if not tx or tx.get("request_only") or tx.get("student_id") != student_id:
            return jsonify({"status": "error", "message": "История платежа не найдена. Старые распределения нельзя отменить автоматически."}), 404
        if tx.get("reversed_at"):
            return jsonify({"status": "error", "message": "Этот платёж уже отменён."}), 409
        schedule = load_json(DATA_FILE)
        students = load_json(STUDENTS_FILE)
        entries = list(student_payment_lessons(schedule, student_id))
        changes = []
        for allocation in tx["allocations"]:
            matches = [(lesson, target) for _date, lesson, target in entries
                       if lesson.get("id") == allocation["lesson_id"]
                       and (lesson.get("lesson_type") == "group") == allocation["is_group"]]
            if len(matches) != 1:
                return jsonify({"status": "error", "message": "Затронутое занятие удалено или изменён его участник. Автоматическая отмена остановлена."}), 409
            lesson, target = matches[0]
            price = member_lesson_price(students, target, lesson) if allocation["is_group"] else get_student_price(students, student_id, lesson.get("price"))
            if comparable_finances(target) != comparable_finances(allocation["after"]) or price != allocation["price"]:
                return jsonify({"status": "error", "message": "Оплата или цена занятия изменились. Сначала отмените более поздние распределения; проверьте стоимость."}), 409
            changes.append((lesson, target, allocation["before"]))
        try:
            with payment_files_transaction():
                for lesson, target, before in changes:
                    for key in FINANCIAL_KEYS:
                        target.pop(key, None)
                    target.update(before)
                    if lesson.get("lesson_type") == "group":
                        lesson["paid"] = bool(lesson.get("group_members")) and all(m.get("paid") or m.get("free") for m in lesson["group_members"])
                now = receipt_now()
                add_receipt_to_book(tx["client_name"], -tx["amount"], tx["receipt_number"], now,
                                    status=f"Отмена общей оплаты · {transaction_id}")
                tx["reversed_at"] = now.isoformat()
                save_json(DATA_FILE, schedule)
                save_json(payments_file(), history)
        except Exception as exc:
            return jsonify({"status": "error", "message": f"Отмена не сохранена, изменения восстановлены: {exc}"}), 500
    return jsonify({"status": "ok"})


def student_payment_candidates(schedule, students, student_id):
    """One chronological queue for custom amounts and 4/8-lesson payments, including debt."""
    candidates = []
    for date_key, lesson, target in student_payment_lessons(schedule, student_id):
        if lesson.get("cancelled") or target.get("free") or target.get("paid") or target.get("paid_via_subscription"):
            continue
        try:
            dt = datetime.datetime.strptime(f"{date_key} {lesson.get('time', '00:00')}", "%Y-%m-%d %H:%M")
        except (TypeError, ValueError):
            continue
        member = target if lesson.get("lesson_type") == "group" else None
        price = member_lesson_price(students, member, lesson) if member is not None else get_student_price(students, student_id, lesson.get("price"))
        if not price or price <= 0:
            continue
        remaining = round(max(0.0, price - max(0.0, normalize_amount(target.get("paid_amount")) or 0.0)), 2)
        if remaining > 0:
            candidates.append((dt, lesson, member, remaining, price))
    candidates.sort(key=lambda row: (row[0], str(row[1].get("id", ""))))
    return candidates


def student_payment_quote(candidates, student_id, count):
    chosen = candidates[:count]
    if len(chosen) < count:
        return None
    allocations = [{"date": dt.strftime("%Y-%m-%d"), "time": dt.strftime("%H:%M"),
                    "id": lesson.get("id"), "is_group": member is not None,
                    "amount": remaining, "price": price,
                    "state": financial_snapshot(member if member is not None else lesson)}
                   for dt, lesson, member, remaining, price in chosen]
    amount = round(sum(row[3] for row in chosen), 2)
    signature = hashlib.sha256(json.dumps(
        {"student_id": student_id, "count": count, "allocations": allocations},
        sort_keys=True, ensure_ascii=False).encode("utf-8")).hexdigest()
    return {"count": count, "amount": amount, "preview_token": signature,
            "lessons": [{key: value for key, value in allocation.items() if key != "state"} for allocation in allocations]}


@flask_app.route("/api/get_student_payment_options", methods=["POST"])
def get_student_payment_options():
    student_id = str((request.get_json() or {}).get("student_id", "")).strip()
    with DATA_LOCK:
        recover_payment_transaction()
        students = load_json(STUDENTS_FILE)
        if student_id not in students:
            return jsonify({"status": "error", "message": "Ученик не найден."}), 404
        candidates = student_payment_candidates(load_json(DATA_FILE), students, student_id)
        options = [quote for count in (4, 8) if (quote := student_payment_quote(candidates, student_id, count))]
    return jsonify({"status": "ok", "options": options, "available_lessons": len(candidates)})


@flask_app.route("/api/apply_student_payment", methods=["POST"])
def apply_student_payment():
    data = request.get_json() or {}
    student_id = str(data.get("student_id", "") or "").strip()
    send_receipt = bool(data.get("send_receipt", True))
    amount = normalize_amount(data.get("amount"))
    # Match the precision used by allocation, receipts and request fingerprints.
    if amount is not None:
        amount = round(amount, 2)
    request_id = data.get("request_id")
    if request_id in (None, ""):
        request_id = ""
    elif not isinstance(request_id, str):
        return jsonify({"status": "error", "message": "Некорректный идентификатор запроса оплаты."}), 400
    else:
        request_id = request_id.strip()
        allowed_request_id_chars = set("abcdefghijklmnopqrstuvwxyzABCDEFGHIJKLMNOPQRSTUVWXYZ0123456789-_.:")
        if not 16 <= len(request_id) <= 128 or any(char not in allowed_request_id_chars for char in request_id):
            return jsonify({"status": "error", "message": "Некорректный идентификатор запроса оплаты."}), 400
    if not student_id:
        return jsonify({"status": "error", "message": "Не указан ученик."}), 400
    if amount is None or amount <= 0:
        return jsonify({"status": "error", "message": "Укажите сумму больше нуля."}), 400

    quick_count = data.get("quick_count")
    if quick_count is not None and (type(quick_count) is not int or quick_count not in (4, 8)):
        return jsonify({"status": "error", "message": "Некорректный вариант оплаты."}), 400
    request_fingerprint = {
        "student_id": student_id,
        "amount": round(amount, 2),
        "quick_count": quick_count,
        "preview_token": str(data.get("preview_token", "") or ""),
        "send_receipt": send_receipt,
    }

    with DATA_LOCK:
        recover_payment_transaction()
        schedule = load_json(DATA_FILE)
        students = load_json(STUDENTS_FILE)
        history = load_json(payments_file(), {})
        if request_id:
            previous = next((tx for tx in history.values()
                             if isinstance(tx, dict) and tx.get("request_id") == request_id), None)
            if previous:
                if previous.get("request_fingerprint") != request_fingerprint:
                    return jsonify({"status": "error", "message": "Этот запрос оплаты уже использован с другими параметрами."}), 409
                if previous.get("reversed_at"):
                    return jsonify({"status": "error", "message": "Этот платёж уже отменён."}), 409
                stored_response = previous.get("idempotency_response") or {
                    "status": "ok",
                    "distributed": previous.get("amount", 0),
                    "unallocated": max(0.0, round((normalize_amount(previous.get("requested_amount")) or 0)
                                                   - (normalize_amount(previous.get("amount")) or 0), 2)),
                    "touched": len(previous.get("allocations") or []),
                    "receipt_number": previous.get("receipt_number"),
                    "receipt_message": "Оплата уже была сохранена ранее.",
                    "transaction_id": previous.get("id"),
                }
                return jsonify({**stored_response, "idempotent_replay": True,
                                **student_lesson_stats(student_id)})
        transaction_id = f"payment_{time.time_ns()}"
        allocations = []
        candidates = student_payment_candidates(schedule, students, student_id)
        if quick_count is not None:
            quote = student_payment_quote(candidates, student_id, quick_count)
            if not quote or quote["preview_token"] != data.get("preview_token") or abs(amount - quote["amount"]) > 0.001:
                return jsonify({"status": "error", "message": "Занятия или оплаты изменились. Откройте оплату заново, чтобы обновить сумму."}), 409
            candidates = candidates[:quick_count]

        left = float(amount)
        distributed = 0.0
        touched = 0
        for _dt, lesson, member, remaining, price in candidates:
            if left <= 0.0001:
                break
            part = min(left, remaining)
            target = member if member is not None else lesson
            before = financial_snapshot(target)
            current = max(0.0, normalize_amount(target.get("paid_amount")) or 0.0)
            target["paid_amount"] = round(current + part, 2)
            left = round(left - part, 2)
            distributed = round(distributed + part, 2)
            touched += 1
            if target["paid_amount"] + 0.0001 >= price:
                target["paid"] = True
                target["paid_amount"] = round(price, 2)
            if member is not None:
                lesson["paid"] = bool(lesson.get("group_members")) and all(bool(m.get("paid")) or bool(m.get("free")) for m in lesson.get("group_members") or [])
            target["allocation_ids"] = [*target.get("allocation_ids", []), transaction_id]
            allocations.append({"lesson_id": lesson.get("id"), "date": _dt.strftime("%Y-%m-%d"),
                                "time": lesson.get("time"), "is_group": member is not None,
                                "price": price, "amount": round(part, 2),
                                "before": before, "after": financial_snapshot(target)})
        settings = load_settings()
        info = get_student_record(students, student_id)
        client_name = str(info.get("name") or "Клиент")
        receipt_path = None
        receipt_number = None
        created_at = None
        if distributed > 0:
            try:
                receipt_path, receipt_number, created_at = generate_receipt_pdf(
                    settings, client_name, distributed, f"student_payment_{student_id}_{time.time_ns()}",
                    service_name_override="Оплата занятий",
                )
            except Exception as exc:
                if receipt_path:
                    delete_receipt_file(receipt_path)
                return jsonify({"status": "error", "message": f"Не удалось сформировать чек: {exc}"}), 500

        try:
            with payment_files_transaction():
                if distributed > 0:
                    add_receipt_to_book(client_name, distributed, receipt_number, created_at, status=f"Оплата занятий · {transaction_id}")
                if distributed > 0 or request_id:
                    response_core = {
                        "status": "ok", "distributed": round(distributed, 2),
                        "unallocated": round(max(0.0, left), 2), "touched": touched,
                        "receipt_number": receipt_number, "receipt_message": "",
                        "transaction_id": transaction_id if distributed > 0 else None,
                    }
                    history[transaction_id] = {"id": transaction_id, "student_id": student_id,
                                               "client_name": client_name, "amount": distributed,
                                               "requested_amount": amount, "receipt_number": receipt_number,
                                               "created_at": (created_at or receipt_now()).isoformat(), "allocations": allocations,
                                               **({"request_only": True} if distributed <= 0 else {}),
                                               **({"request_id": request_id,
                                                   "request_fingerprint": request_fingerprint,
                                                   "idempotency_response": response_core} if request_id else {})}
                    save_json(payments_file(), history)
                save_json(DATA_FILE, schedule)
        except Exception as exc:
            if receipt_path:
                delete_receipt_file(receipt_path)
            return jsonify({"status": "error", "message": f"Оплата не сохранена, изменения отменены: {exc}"}), 500

        request_user = getattr(g, "telegram_user", {}) or {}
        teacher_target = str(getattr(g, "teacher_id", "") or request_user.get("id", "")).strip()
        teacher_chat_id = numeric_telegram_chat_id(teacher_target) if settings.get("default_send_receipt_copy", True) else None
        parent_contacts = info.get("contacts") or {}
        parent_chat_id = numeric_telegram_chat_id(parent_contacts.get("tg") if isinstance(parent_contacts, dict) else "") if send_receipt else None

    messages = []
    try:
        if receipt_path:
            if send_receipt:
                if parent_chat_id is None:
                    messages.append("Родителю чек не отправлен: нет числового Telegram ID.")
                else:
                    ok, error = send_receipt_from_flask(parent_chat_id, receipt_path, f"Чек: оплата занятий · {distributed:.2f} руб. · № {receipt_number}")
                    messages.append("Чек отправлен родителю." if ok else f"Родителю чек отправить не удалось: {error}")
            if settings.get("default_send_receipt_copy", True):
                if teacher_chat_id is not None:
                    ok, error = send_receipt_from_flask(teacher_chat_id, receipt_path, f"Копия чека: оплата занятий · {distributed:.2f} руб. · № {receipt_number}")
                    messages.append("Копия чека отправлена вам." if ok else f"Копию вам отправить не удалось: {error}")
    finally:
        if receipt_path:
            delete_receipt_file(receipt_path)
    return jsonify({
        "status": "ok",
        "distributed": round(distributed, 2),
        "unallocated": round(max(0.0, left), 2),
        "touched": touched,
        "receipt_number": receipt_number,
        "receipt_message": " ".join(messages),
        "transaction_id": transaction_id if distributed > 0 else None,
        **student_lesson_stats(student_id),
    })


@flask_app.route("/api/move_lesson", methods=["POST"])
def move_lesson():
    data = request.get_json() or {}
    old_date = data.get("old_date")
    lesson_id = data.get("id")
    new_date = data.get("new_date")
    new_time = data.get("new_time")
    action_type = data.get("action_type", "move_once")

    try:
        old_dt = datetime.datetime.strptime(old_date, "%Y-%m-%d")
        new_dt = datetime.datetime.strptime(new_date, "%Y-%m-%d")
        datetime.datetime.strptime(new_time, "%H:%M")
    except (TypeError, ValueError):
        return jsonify({"status": "error", "message": "Некорректная дата или время переноса."}), 400

    with DATA_LOCK:
        schedule = load_json(DATA_FILE)
        before = copy.deepcopy(schedule)
        if old_date not in schedule:
            return jsonify({"status": "error", "message": "Исходная дата не найдена."}), 404

        target_lesson = next((lesson.copy() for lesson in schedule[old_date] if lesson.get("id") == lesson_id), None)
        if not target_lesson:
            return jsonify({"status": "error", "message": "Занятие не найдено."}), 404

        if action_type == "copy":
            new_lesson = copy.deepcopy(target_lesson)
            for target in [new_lesson, *new_lesson.get("group_members", [])]:
                for key in (*FINANCIAL_KEYS, "receipt_path", "receipt_book_error"):
                    target.pop(key, None)
                if not is_personal_event(new_lesson):
                    target["paid"] = False
            new_lesson["id"] = next_lesson_id()
            new_lesson["time"] = new_time
            new_lesson.pop("series_id", None)
            new_lesson.pop("reminder_sent_for", None)
            schedule.setdefault(new_date, []).append(new_lesson)

        elif action_type == "move_all":
            series_key = current_series_key(target_lesson)
            day_delta = (new_dt - old_dt).days
            moves = []

            for date_str in sorted(list(schedule.keys())):
                date_obj = datetime.datetime.strptime(date_str, "%Y-%m-%d")
                if date_obj < old_dt:
                    continue

                remaining = []
                for lesson in schedule[date_str]:
                    if same_series(lesson, series_key):
                        moved = lesson.copy()
                        moved["time"] = new_time
                        moved.pop("reminder_sent_for", None)
                        shifted_date = (date_obj + datetime.timedelta(days=day_delta)).strftime("%Y-%m-%d")
                        moves.append((shifted_date, moved))
                    else:
                        remaining.append(lesson)
                schedule[date_str] = remaining

            for date_str in [key for key in list(schedule.keys()) if not schedule[key]]:
                del schedule[date_str]
            for shifted_date, moved in moves:
                schedule.setdefault(shifted_date, []).append(moved)

        else:
            schedule[old_date] = [lesson for lesson in schedule[old_date] if lesson.get("id") != lesson_id]
            if not schedule[old_date]:
                del schedule[old_date]
            target_lesson["time"] = new_time
            target_lesson.pop("reminder_sent_for", None)
            schedule.setdefault(new_date, []).append(target_lesson)

        save_json(DATA_FILE, schedule)
        undo_token = CALENDAR_UNDO.remember(current_teacher_id(), before, schedule)
    return jsonify({"status": "ok", "undo_token": undo_token})


@flask_app.route("/api/delete_lesson", methods=["POST"])
def delete_lesson():
    data = request.get_json() or {}
    date = data.get("date")
    lesson_id = data.get("id")
    delete_all = bool(data.get("delete_all", False))

    with DATA_LOCK:
        schedule = load_json(DATA_FILE)
        before = copy.deepcopy(schedule)
        if date not in schedule or not lesson_id:
            return jsonify({"status": "error", "message": "Занятие не найдено."}), 404

        target_lesson = next((lesson for lesson in schedule[date] if lesson.get("id") == lesson_id), None)
        if not target_lesson:
            return jsonify({"status": "error", "message": "Занятие не найдено."}), 404

        if not delete_all and lesson_has_payment_link(target_lesson):
            return jsonify({"status": "error", "message": "Сначала отмените оплату или используйте отмену занятия вместо удаления."}), 409

        if delete_all:
            series_key = current_series_key(target_lesson)
            start_dt = datetime.datetime.strptime(date, "%Y-%m-%d")
            affected = [lesson for date_str, lessons in schedule.items()
                        if datetime.datetime.strptime(date_str, "%Y-%m-%d") >= start_dt
                        for lesson in lessons if same_series(lesson, series_key)]
            if any(lesson_has_payment_link(lesson) for lesson in affected):
                return jsonify({"status": "error", "message": "В серии есть оплаченные занятия. Сначала отмените их оплату или используйте отмену занятия вместо удаления."}), 409
            for date_str in list(schedule.keys()):
                date_obj = datetime.datetime.strptime(date_str, "%Y-%m-%d")
                if date_obj >= start_dt:
                    schedule[date_str] = [lesson for lesson in schedule[date_str] if not same_series(lesson, series_key)]
                    if not schedule[date_str]:
                        del schedule[date_str]
        else:
            schedule[date] = [lesson for lesson in schedule[date] if lesson.get("id") != lesson_id]
            if not schedule[date]:
                del schedule[date]

        save_json(DATA_FILE, schedule)
        undo_token = CALENDAR_UNDO.remember(current_teacher_id(), before, schedule)
    return jsonify({"status": "ok", "undo_token": undo_token})


@flask_app.route("/api/undo_calendar_action", methods=["POST"])
def undo_calendar_action():
    data = request.get_json(silent=True)
    if not isinstance(data, dict):
        return jsonify({"status": "error", "code": "invalid_request"}), 400
    token = data.get("undo_token")
    owner = current_teacher_id()
    with DATA_LOCK:
        try:
            restored = CALENDAR_UNDO.prepare(owner, token, load_json(DATA_FILE))
        except ValueError as error:
            return jsonify({"status": "error", "code": str(error)}), 409
        save_json(DATA_FILE, restored)
        CALENDAR_UNDO.consume(owner, token)
    return jsonify({"status": "ok"})


async def reminder_worker(application: Application):
    from personal_notifications import send_scheduled_notifications
    try:
        tz = pytz.timezone(TIMEZONE_NAME)
    except pytz.UnknownTimeZoneError:
        tz = pytz.UTC
    while True:
        try:
            await asyncio.to_thread(send_scheduled_notifications, sys.modules[__name__], datetime.datetime.now(tz))
        except Exception:
            print("Scheduled notification check failed")
        await asyncio.sleep(60)


async def post_init(application: Application):
    global BOT_APPLICATION, BOT_LOOP, REMINDER_TASK
    BOT_APPLICATION = application
    BOT_LOOP = asyncio.get_running_loop()
    # post_init выполняется до перехода Application в running-state, поэтому
    # Application.create_task() здесь создаёт предупреждение PTB. Храним обычную
    # asyncio-задачу и явно завершаем её в post_stop.
    REMINDER_TASK = asyncio.create_task(reminder_worker(application), name="schedule-reminder-worker")


async def post_stop(application: Application):
    global REMINDER_TASK, BOT_APPLICATION, BOT_LOOP
    task = REMINDER_TASK
    REMINDER_TASK = None
    if task is not None and not task.done():
        task.cancel()
        try:
            await task
        except asyncio.CancelledError:
            pass
    BOT_APPLICATION = None
    BOT_LOOP = None


async def start(update: Update, context: ContextTypes.DEFAULT_TYPE):
    u = update.effective_user
    from invitation_channels import accept_main, recipient_only
    raw = context.args[0] if context.args else ''
    if raw.startswith('join_'):
        reply = await asyncio.to_thread(accept_main, sys.modules[__name__], raw, u.to_dict(),
                                        update.effective_chat.to_dict(), update.update_id)
        if reply:
            await update.message.reply_text(reply)
        return
    if recipient_only(sys.modules[__name__], str(u.id)):
        await update.message.reply_text('Здесь будут сообщения о занятиях. Для подключения к другому преподавателю откройте его приглашение.')
        return
    # Ordinary entry registers a teacher; invitation visitors were handled above.
    ensure_teacher_registered(str(u.id), {
        "id": str(u.id),
        "first_name": u.first_name or "",
        "last_name": u.last_name or "",
        "username": u.username or "",
    })

    with teacher_scope(str(u.id)):
        language = load_settings().get("language", "ru")
    ready_text = "TEMLI is ready." if language == "en" else "TEMLI готов к работе."
    open_text = "Open TEMLI" if language == "en" else "Открыть TEMLI"
    await update.message.reply_text(
        ready_text,
        reply_markup=InlineKeyboardMarkup([[
            InlineKeyboardButton(open_text, web_app=WebAppInfo(url=WEBAPP_URL))
        ]]),
    )


from personal_bots import register_routes as register_personal_bot_routes
register_personal_bot_routes(sys.modules[__name__])

if __name__ == "__main__":
    from production_server import run
    run(sys.modules[__name__])
